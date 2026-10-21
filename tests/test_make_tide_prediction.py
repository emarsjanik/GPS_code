"""
tests/test_make_tide_prediction.py

The opt-in generator of the predicted-tide spreadsheet the public
7-day plot draws (analysis_tools/make_tide_prediction.py).

The tide is a small synthetic set of constituents with Marconi-like
amplitudes (M2, S2, N2, K1, O1), so every expected answer is known
exactly:
  * the file: the same layout as the hand-made spreadsheet, read the
    way plot_7day.load_tide and station_health read it; 10-minute UTC
    rows over the rolling 30 days; a failed or foreign file is never
    overwritten;
  * NOAA CO-OPS: requests in UTC (time_zone=gmt) and MSL, the 6-minute
    fallback interpolated, a short answer refused, the subordinate
    corrections applied the right way round (a fake server, no
    network);
  * pyTMD: the keywords for pyTMD 2 and 3, times handed over as UTC,
    the land fill (a fake pyTMD module); and, when pyTMD is installed,
    a real run on a synthetic EOT20-format model, checked against an
    independent M2 calculation;
  * validation: known lags (one 10-minute row, a time-zone error, a
    harbour delay, a fraction of a minute), the height ratio and the
    mean offset come back; against NOAA only the difference between
    two spreadsheets' lags is offered as a shift;
  * history: kept by default, started from the spreadsheet
    tide_model_file names, and read back by filter_month.py across a
    month end.

Run with:
    python3 -m unittest discover -s tests -p "test_make_tide_prediction.py" -v
"""

from __future__ import annotations

import contextlib
import importlib.util
import io
import json
import math
import sys
import tempfile
import types
import unittest
import unittest.mock
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import numpy as np

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "analysis_tools"))
sys.path.insert(0, str(ROOT / "diagnostics"))
sys.path.insert(0, str(ROOT / "maintenance"))
import make_tide_prediction as mtp  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

try:                                  # needs matplotlib
    import plot_7day
except Exception:  # pragma: no cover
    plot_7day = None

LAT, LON = 41.8928, -69.9633          # the station (STATION_JSON_REFERENCE.md example)
NOW = datetime(2026, 10, 20, 13, 47)  # the fixed clock of the generation tests
T0 = datetime(2026, 1, 1)

# (constituent, speed in degrees per hour, amplitude m, phase deg):
# Marconi-like amplitudes; the phases are arbitrary.
SYNTH = (
    ("M2", 28.9841042, 1.13, 100.0),
    ("S2", 30.0000000, 0.19, 130.0),
    ("N2", 28.4397295, 0.26, 80.0),
    ("K1", 15.0410686, 0.13, 200.0),
    ("O1", 13.9430356, 0.11, 190.0),
)


def synth_hours(hours) -> np.ndarray:
    hours = np.asarray(hours, dtype=float)
    h = np.zeros_like(hours)
    for _, speed, amp, phase in SYNTH:
        h += amp * np.cos(np.radians(speed * hours - phase))
    return h


def synth(times, late_minutes: float = 0.0) -> np.ndarray:
    """The synthetic tide at naive-UTC `times`, as a file whose clock
    runs `late_minutes` late would hold it: synth(t, L) = synth(t - L)."""
    return synth_hours([(t - T0).total_seconds() / 3600 - late_minutes / 60 for t in times])


def synth_unix(secs) -> np.ndarray:
    """The synthetic tide at UTC seconds since 1970."""
    return synth_hours((np.asarray(secs, dtype=float)
                        - (T0 - datetime(1970, 1, 1)).total_seconds()) / 3600)


def ten_minute(start: datetime, days: float):
    return [start + timedelta(minutes=10 * i) for i in range(int(days * 144))]


# ---------------------------------------------------------------------
# A stand-in for the NOAA CO-OPS API
# ---------------------------------------------------------------------

class _Answer:
    def __init__(self, payload) -> None:
        self._body = json.dumps(payload).encode()

    def read(self) -> bytes:
        return self._body


class FakeNoaa:
    """Serves the synthetic tide as CO-OPS predictions: whole UTC days
    begin_date..end_date at the requested interval, 3 decimals, the
    way the real API answers. `refuse` lists intervals answered with an
    error; `short` drops the last day of every answer; `fail` raises;
    `drop` lists times left out. The gauge's tide is the synthetic one
    `late_minutes` later and `ratio` times as large."""

    def __init__(self, refuse=(), short: bool = False, fail: bool = False, drop=(),
                 late_minutes: float = 0.0, ratio: float = 1.0) -> None:
        self.refuse, self.short, self.fail = set(refuse), short, fail
        self.drop, self.late_minutes, self.ratio = set(drop), late_minutes, ratio
        self.urls = []

    def __call__(self, url, timeout=None):
        self.urls.append(url)
        if self.fail:
            raise OSError("network unreachable")
        q = {k: v[0] for k, v in parse_qs(urlparse(url).query).items()}
        if q["interval"] in self.refuse:
            return _Answer({"error": {"message": "Wrong Interval"}})
        begin = datetime.strptime(q["begin_date"], "%Y%m%d")
        end = datetime.strptime(q["end_date"], "%Y%m%d") + timedelta(days=1)
        if self.short:
            end -= timedelta(days=1)
        step = timedelta(minutes=int(q["interval"]))
        times, t = [], begin
        while t < end:
            if t not in self.drop:
                times.append(t)
            t += step
        values = self.ratio * synth(times, self.late_minutes)
        return _Answer({"predictions": [{"t": t.strftime("%Y-%m-%d %H:%M"), "v": f"{v:.3f}"}
                                        for t, v in zip(times, values)]})


# ---------------------------------------------------------------------
# Stand-ins for pyTMD 2 and 3 (same function, different keywords)
# ---------------------------------------------------------------------

def fake_pytmd(version: int, calls: list, land_until_extrapolated: bool = False,
               all_land: bool = False):
    """sys.modules entries for a pyTMD whose tide_elevations returns the
    synthetic tide at the seconds it is given. The point counts as land
    (NaN) when all_land, or until asked to extrapolate."""

    def answer(secs, extrapolate):
        values = synth_unix(secs)
        if all_land or (land_until_extrapolated and not extrapolate):
            values = np.full(len(values), np.nan)
        return values

    if version == 3:
        def tide_elevations(x, y, delta_time, directory=None, model=None, crop=False,
                            bounds=None, buffer=0, crs=4326, epoch=(2000, 1, 1, 0, 0, 0),
                            type="trajectory", standard="UTC", method="linear",
                            extrapolate=False, cutoff=10.0, **kwargs):
            calls.append(dict(kwargs, x=x, y=y, directory=directory, model=model, crop=crop,
                              crs=crs, epoch=epoch, type=type, standard=standard,
                              method=method, extrapolate=extrapolate, cutoff=cutoff))
            return types.SimpleNamespace(values=answer(delta_time, extrapolate))
    else:
        def tide_elevations(x, y, delta_time, DIRECTORY=None, MODEL=None, CROP=False,
                            EPSG=4326, EPOCH=(2000, 1, 1, 0, 0, 0), TYPE="drift",
                            TIME="UTC", METHOD="spline", EXTRAPOLATE=False, CUTOFF=10.0,
                            **kwargs):
            calls.append(dict(kwargs, x=x, y=y, DIRECTORY=DIRECTORY, MODEL=MODEL, CROP=CROP,
                              EPSG=EPSG, EPOCH=EPOCH, TYPE=TYPE, TIME=TIME, METHOD=METHOD,
                              EXTRAPOLATE=EXTRAPOLATE, CUTOFF=CUTOFF))
            return np.ma.masked_invalid(answer(delta_time, EXTRAPOLATE))

    pkg = types.ModuleType("pyTMD")
    pkg.__version__ = f"{version}.0-fake"
    compute = types.ModuleType("pyTMD.compute")
    compute.tide_elevations = tide_elevations
    pkg.compute = compute
    return {"pyTMD": pkg, "pyTMD.compute": compute}


NO_PYTMD = {"pyTMD": None, "pyTMD.compute": None}     # import raises ImportError


def opts(**over) -> dict:
    o = {"source": "auto", "lat": LAT, "lon": LON, "pytmd_directory": None,
         "model": "EOT20", "extrapolate_km": 10.0, "noaa_station": "8447435",
         "noaa_time_offset_minutes": 0.0, "noaa_height_ratio": 1.0,
         "keep_history": False, "value_column": None, "time_column": "time"}
    o.update(over)
    return o


class TestWindow(unittest.TestCase):

    def test_rolling_window_runs_30_days_from_10_days_back(self) -> None:
        start, end = mtp.rolling_window(NOW)
        self.assertEqual(start, datetime(2026, 10, 10))
        self.assertEqual(end, datetime(2026, 11, 9))

    def test_times_fall_on_whole_10_minute_marks(self) -> None:
        times = mtp.grid_times(datetime(2026, 10, 10, 0, 3), datetime(2026, 10, 11))
        self.assertEqual(times[0], datetime(2026, 10, 10, 0, 10))
        self.assertEqual(times[-1], datetime(2026, 10, 10, 23, 50))
        self.assertTrue(all(t.minute % 10 == 0 and t.second == 0 for t in times))
        self.assertEqual(len(mtp.grid_times(*mtp.rolling_window(NOW))), 30 * 144)


class TestNoaa(unittest.TestCase):

    def setUp(self) -> None:
        self.times = mtp.grid_times(datetime(2026, 10, 10), datetime(2026, 11, 9))

    def test_ten_minute_predictions_are_requested_in_utc_and_msl(self) -> None:
        server = FakeNoaa()
        got = mtp.noaa_prediction(self.times, urlopen=server)
        np.testing.assert_allclose(got, synth(self.times), atol=0.0005)
        for url in server.urls:
            q = parse_qs(urlparse(url).query)
            self.assertEqual(q["time_zone"], ["gmt"])
            self.assertEqual(q["datum"], ["MSL"])
            self.assertEqual(q["units"], ["metric"])
            self.assertEqual(q["interval"], ["10"])
            self.assertEqual(q["station"], ["8447435"])
        # 31 UTC days (one after the window) in 30-day requests
        self.assertEqual(len(server.urls), 2)

    def test_six_minute_fallback_is_interpolated_to_the_grid(self) -> None:
        server = FakeNoaa(refuse={"10"})
        with contextlib.redirect_stdout(io.StringIO()) as out:
            got = mtp.noaa_prediction(self.times, urlopen=server)
        self.assertIn("6-minute", out.getvalue())
        # linear interpolation of 6-minute samples: well under 2 mm here
        self.assertLess(np.max(np.abs(got - synth(self.times))), 0.002)

    def test_short_answer_raises_instead_of_leaving_a_hole(self) -> None:
        with contextlib.redirect_stdout(io.StringIO()):
            with self.assertRaises(RuntimeError):
                mtp.noaa_prediction(self.times, urlopen=FakeNoaa(short=True))

    def test_one_missing_sample_raises(self) -> None:
        server = FakeNoaa(drop={datetime(2026, 10, 20, 12, 10)})
        with self.assertRaisesRegex(RuntimeError, "missing at 1 time"):
            mtp.noaa_prediction(self.times, urlopen=server)

    def test_subordinate_corrections_shift_later_and_scale(self) -> None:
        got = mtp.noaa_prediction(self.times, time_offset_minutes=30, height_ratio=0.9,
                                  urlopen=FakeNoaa())
        # the site's tide is the station's, 30 min LATER and 0.9 times as large
        np.testing.assert_allclose(got, 0.9 * synth(self.times, late_minutes=30), atol=0.0005)


class TestPytmdKeywords(unittest.TestCase):

    def setUp(self) -> None:
        self.times = ten_minute(datetime(2026, 10, 1), 2)

    def test_pytmd3_gets_utc_seconds_and_lower_case_keywords(self) -> None:
        calls = []
        with unittest.mock.patch.dict(sys.modules, fake_pytmd(3, calls)):
            got = mtp.pytmd_prediction(self.times, LAT, LON, directory="/models")
        np.testing.assert_allclose(got, synth(self.times), atol=1e-9)
        (kw,) = calls
        self.assertEqual((kw["standard"], kw["epoch"], kw["type"]),
                         ("UTC", (1970, 1, 1, 0, 0, 0), "trajectory"))
        self.assertEqual((kw["model"], kw["directory"], kw["crs"]), ("EOT20", "/models", 4326))
        self.assertEqual(kw["method"], "linear")
        self.assertTrue(kw["infer_minor"])
        self.assertFalse(kw["extrapolate"])
        # cropping needs dask in pyTMD 3
        self.assertEqual(kw["crop"], importlib.util.find_spec("dask") is not None)
        np.testing.assert_array_equal(kw["x"], LON)
        np.testing.assert_array_equal(kw["y"], LAT)

    def test_pytmd2_gets_upper_case_keywords_and_no_crop(self) -> None:
        calls = []
        with unittest.mock.patch.dict(sys.modules, fake_pytmd(2, calls)):
            got = mtp.pytmd_prediction(self.times, LAT, LON)
        np.testing.assert_allclose(got, synth(self.times), atol=1e-9)
        (kw,) = calls
        self.assertEqual((kw["TIME"], kw["EPOCH"], kw["TYPE"]),
                         ("UTC", (1970, 1, 1, 0, 0, 0), "drift"))
        self.assertEqual((kw["MODEL"], kw["METHOD"], kw["CROP"]), ("EOT20", "linear", False))
        self.assertIsNone(kw["DIRECTORY"])

    def test_land_cell_is_filled_from_the_ocean_nearby(self) -> None:
        calls = []
        with unittest.mock.patch.dict(sys.modules,
                                      fake_pytmd(3, calls, land_until_extrapolated=True)):
            got = mtp.pytmd_prediction(self.times, LAT, LON, extrapolate_km=5.0)
        self.assertEqual([c["extrapolate"] for c in calls], [False, True])
        self.assertEqual(calls[1]["cutoff"], 5.0)
        np.testing.assert_allclose(got, synth(self.times), atol=1e-9)

    def test_land_fill_is_said_and_recorded(self) -> None:
        with unittest.mock.patch.dict(sys.modules,
                                      fake_pytmd(3, [], land_until_extrapolated=True)), \
                contextlib.redirect_stdout(io.StringIO()) as out:
            _, about = mtp.predict(self.times, opts(source="pytmd"))
        self.assertIn("station cell is land", out.getvalue())
        self.assertIn("nearest ocean node within 10 km", about["land_fill"])
        with unittest.mock.patch.dict(sys.modules, fake_pytmd(3, [])):
            _, about = mtp.predict(self.times, opts(source="pytmd"))
        self.assertNotIn("land_fill", about)

    def test_no_ocean_within_reach_raises(self) -> None:
        with unittest.mock.patch.dict(sys.modules, fake_pytmd(3, [], all_land=True)):
            with self.assertRaisesRegex(RuntimeError, "offshore"):
                mtp.pytmd_prediction(self.times, LAT, LON)

    def test_auto_falls_back_to_noaa_and_says_why(self) -> None:
        with unittest.mock.patch.dict(sys.modules, NO_PYTMD), \
                unittest.mock.patch("urllib.request.urlopen", FakeNoaa()), \
                contextlib.redirect_stdout(io.StringIO()) as out:
            values, about = mtp.predict(self.times, opts(source="auto"))
        self.assertIn("not usable here", out.getvalue())
        self.assertTrue(about["source_key"].startswith("noaa:8447435"))
        self.assertEqual(about["value_column"], "NOAA8447435_heightm")
        np.testing.assert_allclose(values, synth(self.times), atol=0.0005)

    def test_pytmd_source_does_not_fall_back(self) -> None:
        with unittest.mock.patch.dict(sys.modules, NO_PYTMD), \
                unittest.mock.patch("urllib.request.urlopen", FakeNoaa()) as server:
            with self.assertRaises(ImportError):
                mtp.predict(self.times, opts(source="pytmd"))
        self.assertEqual(server.urls, [])

    def test_eot20_column_matches_the_spreadsheet(self) -> None:
        with unittest.mock.patch.dict(sys.modules, fake_pytmd(3, [])):
            _, about = mtp.predict(self.times, opts(source="pytmd"))
        self.assertEqual(about["value_column"], "EOT20_heightm")
        self.assertEqual(about["source"], "pyTMD 3.0-fake, EOT20 ocean tide")


class TestValidation(unittest.TestCase):
    """compare_series must give back the lag, ratio and offset that were
    put into a reference."""

    def setUp(self) -> None:
        self.gen_t = ten_minute(datetime(2026, 7, 14), 50)
        self.gen_v = synth(self.gen_t)
        self.ref_t = ten_minute(datetime(2026, 7, 15), 47)

    def compare(self, ref_v):
        return mtp.compare_series(self.ref_t, ref_v, self.gen_t, self.gen_v)

    def test_identical_series_agree(self) -> None:
        res = self.compare(synth(self.ref_t))
        self.assertLess(res["rms"], 1e-9)
        self.assertAlmostEqual(res["lag_min"], 0.0, delta=0.05)
        self.assertAlmostEqual(res["ratio"], 1.0, places=6)
        self.assertEqual(res["off_grid"], 0)
        self.assertEqual(mtp.read_lag(res["lag_min"]), "the two agree in time")

    def test_one_ten_minute_row_is_found(self) -> None:
        # the older spreadsheet's clock runs 10 min early (merge_tide_files.py)
        res = self.compare(synth(self.ref_t, late_minutes=-10))
        self.assertAlmostEqual(res["lag_min"], -10.0, delta=0.2)
        self.assertLess(res["rms_at_lag"], 0.002)
        self.assertGreater(res["rms"], 0.05)       # ~10 cm at peak flow
        self.assertIn("one 10-minute row", mtp.read_lag(res["lag_min"]))
        text = mtp.format_report("old.xlsx", "EOT20_heightm", "test", res)
        self.assertIn("--shift-old-minutes +10", text)

    def test_eastern_time_written_as_utc_is_found(self) -> None:
        # EDT labels are 4 h behind UTC: the file's clock runs 4 h early
        res = self.compare(synth(self.ref_t, late_minutes=-240))
        self.assertAlmostEqual(res["lag_min"], -240.0, delta=0.5)
        self.assertIn("EDT", mtp.read_lag(res["lag_min"]))

    def test_harbour_delay_range_and_offset(self) -> None:
        ref = 0.8 * synth(self.ref_t, late_minutes=37) + 0.05
        res = self.compare(ref)
        self.assertAlmostEqual(res["lag_min"], 37.0, delta=0.5)
        self.assertAlmostEqual(res["ratio"], 0.8, delta=0.005)
        self.assertAlmostEqual(res["mean"], 0.05, delta=0.005)
        self.assertIn("not a clock error", mtp.read_lag(res["lag_min"]))

    def test_lag_is_found_to_a_fraction_of_a_minute(self) -> None:
        for late in (-9.6, 2.7):
            res = self.compare(synth(self.ref_t, late_minutes=late))
            self.assertAlmostEqual(res["lag_min"], late, delta=0.05)

    def test_a_different_range_does_not_move_the_lag(self) -> None:
        # on a whole row, where the prediction needs no interpolation
        for ratio in (0.8, 1.25):
            res = self.compare(ratio * synth(self.ref_t, late_minutes=-10))
            self.assertAlmostEqual(res["lag_min"], -10.0, delta=0.05)

    def test_every_lag_is_scored_on_rows_the_prediction_covers_at_all_lags(self) -> None:
        # 1.4 days of overlap: enough to compare, but only ~0.9 days
        # have the prediction on both sides at +/-6 h, too few to scan
        gen_t = ten_minute(datetime(2026, 7, 20), 1.4)
        res = mtp.compare_series(gen_t, synth(gen_t), gen_t, synth(gen_t))
        self.assertIsNotNone(res)
        self.assertIsNone(res["lag_min"])

    def test_positive_time_zone_lag_is_not_eastern_time_written_as_utc(self) -> None:
        self.assertIn("written as UTC", mtp.read_lag(-240.0))
        self.assertIn("4 h late", mtp.read_lag(240.0))
        self.assertNotIn("written as UTC", mtp.read_lag(300.0))

    def test_against_noaa_no_shift_is_offered_from_one_lag(self) -> None:
        res = self.compare(0.8 * synth(self.ref_t, late_minutes=-10))
        text = mtp.format_report("old.xlsx", "EOT20_heightm", "NOAA", res, kind="noaa")
        self.assertNotIn("--shift-old-minutes", text)
        self.assertNotIn("add", text)
        self.assertIn("not a clock error by itself", text)
        text = mtp.format_report("old.xlsx", "EOT20_heightm", "EOT20", res, kind="pytmd")
        self.assertIn("--shift-old-minutes +10", text)

    def test_shift_is_offered_only_for_a_clock_error(self) -> None:
        res = self.compare(synth(self.ref_t, late_minutes=37))
        text = mtp.format_report("x.xlsx", "EOT20_heightm", "EOT20", res, kind="pytmd")
        self.assertNotIn("--shift-old-minutes", text)

    def test_gap_in_the_prediction_is_not_bridged(self) -> None:
        keep = [not (datetime(2026, 8, 1) <= t < datetime(2026, 8, 2)) for t in self.gen_t]
        gen_t = [t for t, k in zip(self.gen_t, keep) if k]
        gen_v = self.gen_v[np.array(keep)]
        res = mtp.compare_series(self.ref_t, synth(self.ref_t), gen_t, gen_v)
        # the day's 144 rows are left out, not interpolated across
        self.assertEqual(res["n"], len(self.ref_t) - 144)
        self.assertLess(res["rms"], 1e-9)

    def test_too_short_an_overlap_gives_nothing(self) -> None:
        ref_t = ten_minute(datetime(2026, 9, 2, 12), 1)
        self.assertIsNone(mtp.compare_series(ref_t, synth(ref_t), self.gen_t, self.gen_v))


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class TestGenerate(unittest.TestCase):

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.dir = Path(self._tmp.name)
        self.out = self.dir / "tide_prediction.xlsx"
        self.cfg = {"latitude": LAT, "longitude": LON}
        p = unittest.mock.patch.object(mtp, "load_station", lambda: self.cfg)
        p.start()
        self.addCleanup(p.stop)
        self.now = NOW
        p = unittest.mock.patch.object(mtp, "utcnow", lambda: self.now)
        p.start()
        self.addCleanup(p.stop)

    def run_main(self, *argv, server=None, modules=None):
        server = server or FakeNoaa()
        with unittest.mock.patch.dict(sys.modules, modules or NO_PYTMD), \
                unittest.mock.patch("urllib.request.urlopen", server), \
                contextlib.redirect_stdout(io.StringIO()) as out:
            code = mtp.main(["--output", str(self.out), *argv])
        return code, out.getvalue()

    def hand_made(self, path: Path) -> None:
        wb = openpyxl.Workbook()
        wb.active.append(["time", "EOT20_heightm"])
        for t in ten_minute(datetime(2026, 9, 1), 2):
            wb.active.append([t, 0.5])
        wb.save(path)

    def test_file_has_the_spreadsheet_layout(self) -> None:
        code, _ = self.run_main("--source", "noaa")
        self.assertEqual(code, 0)
        wb = openpyxl.load_workbook(self.out, data_only=True)
        ws = wb[wb.sheetnames[0]]
        self.assertEqual([c.value for c in ws[1]], ["time", "NOAA8447435_heightm"])
        self.assertEqual(wb.sheetnames[1], "about")
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        wb.close()
        times = [r[0] for r in rows]
        self.assertEqual(len(rows), 30 * 144)
        self.assertEqual((times[0], times[-1]),
                         (datetime(2026, 10, 10), datetime(2026, 11, 8, 23, 50)))
        self.assertTrue(all(b - a == timedelta(minutes=10) for a, b in zip(times, times[1:])))
        np.testing.assert_allclose([r[1] for r in rows], synth(times), atol=0.0005)
        times, values, about, column = mtp.read_prediction(self.out)
        self.assertEqual(column, "NOAA8447435_heightm")
        self.assertTrue(mtp.is_generated(about))
        self.assertIn("UTC", about["time"])

    @unittest.skipIf(plot_7day is None, "plot_7day needs matplotlib")
    def test_the_7_day_plot_reads_it_unchanged(self) -> None:
        self.run_main("--source", "noaa")
        times, values = plot_7day.load_tide(self.out, "time", "NOAA8447435_heightm")
        self.assertEqual(len(times), 30 * 144)
        np.testing.assert_allclose(values, synth(times), atol=0.0005)

    def test_the_health_check_reads_its_span(self) -> None:
        import station_health as sh
        self.run_main("--source", "noaa")
        want = (datetime(2026, 10, 10), datetime(2026, 11, 8, 23, 50))
        self.assertEqual(sh._xlsx_column_range(self.out, "time"), want)
        got = sh._xlsx_column_range_stdlib(self.out, "time")
        self.assertTrue(all(abs((g - w).total_seconds()) < 1 for g, w in zip(got, want)))

    def test_hand_made_spreadsheet_is_never_overwritten(self) -> None:
        self.hand_made(self.out)
        before = self.out.read_bytes()
        code, out = self.run_main("--source", "noaa")
        self.assertEqual(code, 2)
        self.assertIn("left alone", out)
        self.assertEqual(self.out.read_bytes(), before)
        code, _ = self.run_main("--source", "noaa", "--force")
        self.assertEqual(code, 0)
        self.assertTrue(mtp.is_generated(mtp.read_prediction(self.out)[2]))

    def test_failed_refresh_leaves_the_previous_file(self) -> None:
        self.run_main("--source", "noaa")
        before = self.out.read_bytes()
        self.now = NOW + timedelta(days=1)
        code, out = self.run_main("--source", "noaa", server=FakeNoaa(fail=True))
        self.assertEqual(code, 1)
        self.assertIn("unchanged", out)
        self.assertEqual(self.out.read_bytes(), before)
        self.assertEqual(sorted(p.name for p in self.dir.iterdir()), [self.out.name])

    def test_keep_history_keeps_earlier_rows_of_the_same_source(self) -> None:
        self.run_main("--source", "noaa")
        self.now = NOW + timedelta(days=3)
        code, out = self.run_main("--source", "noaa", "--keep-history")
        self.assertEqual(code, 0)
        times, values, about, _ = mtp.read_prediction(self.out)
        self.assertEqual((times[0], times[-1]),
                         (datetime(2026, 10, 10), datetime(2026, 11, 11, 23, 50)))
        self.assertTrue(all(b - a == timedelta(minutes=10) for a, b in zip(times, times[1:])))
        self.assertEqual(about["history_from"], "2026-10-10 00:00 UTC")
        self.assertIn("kept 432 earlier row(s)", out)

    def test_history_is_kept_by_default_and_can_be_switched_off(self) -> None:
        self.run_main("--source", "noaa")
        self.now = NOW + timedelta(days=3)
        code, out = self.run_main("--source", "noaa")
        self.assertEqual(code, 0)
        self.assertIn("kept 432 earlier row(s)", out)
        self.assertEqual(mtp.read_prediction(self.out)[0][0], datetime(2026, 10, 10))
        self.now = NOW + timedelta(days=4)
        code, out = self.run_main("--source", "noaa", "--no-keep-history")
        self.assertEqual(code, 0)
        self.assertEqual(mtp.read_prediction(self.out)[0][0], datetime(2026, 10, 14))

    def test_keep_history_must_be_a_boolean(self) -> None:
        self.cfg["tide_prediction"] = {"keep_history": "false"}
        with self.assertRaises(SystemExit):
            self.run_main("--source", "noaa")

    def test_first_run_starts_from_the_spreadsheet_and_monthly_csvs_keep_their_tide(self) -> None:
        # The hand-made spreadsheet the plot reads ends on 30 September;
        # the switch comes in October. Both months must keep every tide
        # value when filter_month.py makes their CSVs again in November.
        import filter_month
        sheet = self.dir / "marconi_tides.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["time", "EOT20_heightm"])
        sheet_t = ten_minute(datetime(2026, 9, 1), 30)
        sheet_v = [round(float(v), 4) for v in synth(sheet_t, late_minutes=-10)]
        for t, v in zip(sheet_t, sheet_v):
            wb.active.append([t, v])
        wb.save(sheet)
        self.cfg.update(tide_model_file=str(sheet), tide_model_value_column="EOT20_heightm")
        code, out = self.run_main("--source", "pytmd", modules=fake_pytmd(3, []))
        self.assertEqual(code, 0)
        self.assertIn(f"took {len(sheet_t)} earlier row(s) from marconi_tides.xlsx", out)
        times, values, about, _ = mtp.read_prediction(self.out)
        self.assertEqual((times[0], times[-1]),
                         (datetime(2026, 9, 1), datetime(2026, 11, 8, 23, 50)))
        self.assertTrue(all(b - a == timedelta(minutes=10) for a, b in zip(times, times[1:])))
        self.assertEqual(list(values[:len(sheet_t)]), sheet_v)        # copied as they were
        np.testing.assert_allclose(values[len(sheet_t):], synth(times[len(sheet_t):]),
                                   atol=0.0001)
        self.assertIn("marconi_tides.xlsx", about["seeded_from"])

        # now the plot reads the generated file; three weeks of nights later
        self.cfg["tide_model_file"] = str(self.out)
        self.now = NOW + timedelta(days=25)
        code, out = self.run_main("--source", "pytmd", modules=fake_pytmd(3, []))
        self.assertEqual(code, 0)
        about = mtp.read_prediction(self.out)[2]
        self.assertIn("marconi_tides.xlsx", about["seeded_from"])
        tide_t, tide_v = filter_month.load_tide(self.out, "time", "EOT20_heightm")
        for month in (9, 10):
            query = [t for t in ten_minute(datetime(2026, month, 1), 31)[::3]
                     if t.month == month]
            got = filter_month.interp_tide(tide_t, tide_v, query)
            self.assertNotIn(None, got)
        got = filter_month.interp_tide(tide_t, tide_v, sheet_t)
        self.assertEqual(got, sheet_v)

    def test_spreadsheet_with_another_column_is_not_copied(self) -> None:
        sheet = self.dir / "noaa_sheet.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["time", "FES_heightm"])
        for t in ten_minute(datetime(2026, 9, 1), 2):
            wb.active.append([t, 0.5])
        wb.save(sheet)
        self.cfg.update(tide_model_file=str(sheet), tide_model_value_column="FES_heightm")
        code, out = self.run_main("--source", "pytmd", modules=fake_pytmd(3, []))
        self.assertEqual(code, 0)
        self.assertIn("no rows taken from noaa_sheet.xlsx", out)
        self.assertEqual(mtp.read_prediction(self.out)[0][0], datetime(2026, 10, 10))

    def test_auto_does_not_rename_the_column_the_plot_reads(self) -> None:
        # tide_model_file is this file and reads EOT20_heightm: with no
        # pyTMD, auto fails rather than write a NOAA column
        self.cfg.update(tide_model_file=str(self.out), tide_model_value_column="EOT20_heightm")
        server = FakeNoaa()
        code, out = self.run_main("--source", "auto", server=server)
        self.assertEqual(code, 1)
        self.assertIn("auto: using pytmd", out)
        self.assertEqual(server.urls, [])
        self.assertFalse(self.out.exists())

    def test_history_from_another_source_is_not_kept(self) -> None:
        self.run_main("--source", "noaa")
        self.now = NOW + timedelta(days=3)
        code, out = self.run_main("--source", "noaa", "--noaa-station", "8443970",
                                  "--keep-history")
        self.assertEqual(code, 0)
        self.assertIn("not kept", out)
        times, _, _, column = mtp.read_prediction(self.out)
        self.assertEqual(column, "NOAA8443970_heightm")
        self.assertEqual(times[0], datetime(2026, 10, 13))

    def test_auto_keeps_the_source_of_its_existing_file(self) -> None:
        # pyTMD missing on the first night: NOAA. Later pyTMD works, but
        # the plot reads NOAA8447435_heightm, so the column must not change.
        self.run_main("--source", "auto")
        calls = []
        self.now = NOW + timedelta(days=1)
        code, out = self.run_main("--source", "auto", modules=fake_pytmd(3, calls))
        self.assertEqual(code, 0)
        self.assertIn("keeping noaa", out)
        self.assertEqual(calls, [])
        self.assertEqual(mtp.read_prediction(self.out)[3], "NOAA8447435_heightm")

    def test_station_json_settings_are_used(self) -> None:
        self.cfg["tide_prediction"] = {"source": "noaa", "file": str(self.out),
                                       "noaa_station": "8443970", "point": [41.9, -69.9]}
        server = FakeNoaa()
        with unittest.mock.patch.dict(sys.modules, NO_PYTMD), \
                unittest.mock.patch("urllib.request.urlopen", server), \
                contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(mtp.main([]), 0)
        self.assertIn("station=8443970", server.urls[0])
        self.assertEqual(mtp.read_prediction(self.out)[3], "NOAA8443970_heightm")

    def test_malformed_point_is_refused(self) -> None:
        self.cfg["tide_prediction"] = {"point": 41.9}
        with self.assertRaises(SystemExit):
            self.run_main("--source", "noaa")

    def write_reference(self, name: str, late_minutes: float) -> Path:
        path = self.dir / name
        wb = openpyxl.Workbook()
        wb.active.append(["time", "EOT20_heightm"])
        times = ten_minute(datetime(2026, 7, 15), 10)
        for t, v in zip(times, synth(times, late_minutes)):
            wb.active.append([t, round(float(v), 4)])
        wb.save(path)
        return path

    def test_validate_against_two_spreadsheets(self) -> None:
        old = self.write_reference("old.xlsx", late_minutes=-10)
        new = self.write_reference("new.xlsx", late_minutes=0)
        gen = self.dir / "gen.xlsx"
        times = ten_minute(datetime(2026, 7, 14), 12)
        mtp.write_prediction(gen, times, synth(times), "time", "EOT20_heightm",
                             {"source": "test", "source_key": "pytmd:EOT20:test"})
        code, out = self.run_main("--validate", "--reference", str(old), str(new),
                                  "--generated", str(gen))
        self.assertEqual(code, 0)
        old_part, new_part = out.split("Spreadsheet   : ")[1:]
        new_part, between = new_part.split("Between the spreadsheets")
        self.assertIn("Best lag      : -10.0 min", old_part)
        self.assertIn("--shift-old-minutes +10", old_part)
        self.assertIn("Best lag      : +0.0 min", new_part)
        self.assertIn("the two agree in time", new_part)
        self.assertIn("new.xlsx lag minus old.xlsx lag: +10.0 min", between)
        self.assertIn("--shift-old-minutes +10", between)
        self.assertFalse(self.out.exists())           # validation writes nothing

    def test_validate_against_noaa_gives_the_shift_between_the_spreadsheets(self) -> None:
        # No pyTMD, so auto uses NOAA, whose gauge runs 37 min late at
        # 0.8 x: each lag holds that delay, only their difference does not
        old = self.write_reference("old.xlsx", late_minutes=-10)
        new = self.write_reference("new.xlsx", late_minutes=0)
        code, out = self.run_main("--validate", "--reference", str(old), str(new),
                                  server=FakeNoaa(late_minutes=37, ratio=0.8))
        self.assertEqual(code, 0)
        self.assertIn("using NOAA CO-OPS", out)
        reports, between = out.split("Between the spreadsheets")
        self.assertIn("Best lag      : -47.0 min", reports)
        self.assertIn("Best lag      : -37.0 min", reports)
        self.assertNotIn("--shift-old-minutes", reports)
        self.assertNotIn("add +", reports)
        self.assertIn("lag minus old.xlsx lag: +10.0 min", between)
        self.assertIn("--old old.xlsx --new new.xlsx --shift-old-minutes +10", between)
        self.assertIn("one 10-minute row", between)

    def test_validate_makes_its_own_prediction(self) -> None:
        ref = self.write_reference("edt.xlsx", late_minutes=-240)
        calls = []
        code, out = self.run_main("--validate", "--source", "pytmd", "--reference", str(ref),
                                  modules=fake_pytmd(3, calls))
        self.assertEqual(code, 0)
        self.assertEqual(len(calls), 1)
        self.assertIn("Best lag      : -240.0 min", out)
        self.assertIn("a time-zone error", out)


def _netcdf_writer():
    if importlib.util.find_spec("netCDF4"):
        return "netCDF4"
    if importlib.util.find_spec("xarray") and importlib.util.find_spec("h5netcdf"):
        return "xarray"
    return None


EOT20_FILES = ("2N2", "J1", "K1", "K2", "M2", "M4", "MF", "MM", "N2", "O1", "P1", "Q1",
               "S1", "S2", "SA", "SSA", "T2")


def independent_m2(t: datetime, amp: float, g: float) -> float:
    """M2 at naive-UTC t from Meeus's mean longitudes and Schureman's
    nodal factors, written here independently of pyTMD: V = 2(T + h - s)
    with T the mean solar hour angle (TT taken as UTC + 69.2 s)."""
    jd = 2440587.5 + ((t - datetime(1970, 1, 1)).total_seconds() + 69.2) / 86400
    c = (jd - 2451545.0) / 36525
    s = 218.3164591 + 481267.88134236 * c
    h = 280.4664567 + 36000.7697489 * c
    node = math.radians(125.0445550 - 1934.1361849 * c)
    ut = t.hour + t.minute / 60 + t.second / 3600
    v = 2 * (15 * ut + 180 + h - s)
    f = 1.0004 - 0.0373 * math.cos(node) + 0.0002 * math.cos(2 * node)
    u = -2.14 * math.sin(node)
    return f * amp * math.cos(math.radians(v + u - g))


@unittest.skipUnless(importlib.util.find_spec("pyTMD") and _netcdf_writer(),
                     "pyTMD (and a netCDF writer) not installed")
class TestRealPytmd(unittest.TestCase):
    """The installed pyTMD on a synthetic EOT20-format model: M2 only,
    1 m with Greenwich phase 100 deg everywhere. Agreement with the
    independent M2 to a few mm, at zero lag, shows the times reach
    pyTMD as UTC and the keywords suit this pyTMD version."""

    def test_m2_only_model_matches_an_independent_calculation(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / "EOT20" / "ocean_tides"
            folder.mkdir(parents=True)
            lon = np.arange(0.0, 360.0 + 1e-9, 1.0)
            lat = np.arange(-90.0, 90.0 + 1e-9, 1.0)
            for name in EOT20_FILES:
                # a zero amplitude would read as land: 0.01 cm elsewhere
                amp = np.full((lat.size, lon.size), 100.0 if name == "M2" else 0.01)
                pha = np.full((lat.size, lon.size), 100.0 if name == "M2" else 10.0)
                self._write(folder / f"{name}_ocean_eot20.nc", lat, lon, amp, pha)
            times = ten_minute(datetime(2026, 10, 1), 2)
            got = mtp.pytmd_prediction(times, LAT, LON, directory=tmp, infer_minor=False)
        want = np.array([independent_m2(t, 1.0, 100.0) for t in times])
        self.assertLess(np.max(np.abs(got - want)), 0.003)
        res = mtp.compare_series(times[36:-36], got[36:-36], times, want, max_lag_minutes=60)
        self.assertAlmostEqual(res["lag_min"], 0.0, delta=0.3)

    @staticmethod
    def _write(path, lat, lon, amp, pha) -> None:
        if _netcdf_writer() == "netCDF4":
            import netCDF4
            with netCDF4.Dataset(path, "w") as ds:
                ds.createDimension("lat", lat.size)
                ds.createDimension("lon", lon.size)
                for name, data, units, dims in (("lat", lat, "degrees_north", ("lat",)),
                                                ("lon", lon, "degrees_east", ("lon",)),
                                                ("amplitude", amp, "cm", ("lat", "lon")),
                                                ("phase", pha, "degrees", ("lat", "lon"))):
                    var = ds.createVariable(name, "f8" if len(dims) == 1 else "f4", dims)
                    var[:] = data
                    var.units = units
        else:
            import xarray as xr
            xr.Dataset({"amplitude": (("lat", "lon"), amp.astype("f4"), {"units": "cm"}),
                        "phase": (("lat", "lon"), pha.astype("f4"), {"units": "degrees"})},
                       coords={"lat": ("lat", lat, {"units": "degrees_north"}),
                               "lon": ("lon", lon, {"units": "degrees_east"})}
                       ).to_netcdf(path, engine="h5netcdf")


if __name__ == "__main__":
    unittest.main()
