"""
tests/test_station_health.py

The health checks that make silent failures visible in the morning
mail: a 7_day_plot.png that stopped being redrawn (process_and_plot.sh
reports success when the plot fails), and a tide prediction
spreadsheet that has run out (the predicted tide and the departure
markers then vanish from the plot without a word).
"""

from __future__ import annotations

import json
import os
import sys
import tempfile
import time
import unittest
import unittest.mock
import zipfile
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "diagnostics"))
import station_health as sh  # noqa: E402
import gnss_record as gr  # noqa: E402

STA = "usgs"
MJD0 = datetime(1858, 11, 17)

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


class HealthTest(unittest.TestCase):

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.proj = Path(self._tmp.name)
        self.files = self.proj / "products" / "refl_code" / "Files" / STA
        self.files.mkdir(parents=True)
        patch = unittest.mock.patch.object(sh, "PROJECT_DIR", self.proj)
        patch.start()
        self.addCleanup(patch.stop)
        sh.findings.clear()

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def findings(self, check: str):
        return [(f[0], f[2]) for f in sh.findings if f[1] == check]

    def write_spline(self, newest: datetime, mtime: float | None = None) -> Path:
        path = self.files / f"{STA}_spline_out.txt"
        lines = ["% fake spline\n"]
        t = newest - timedelta(days=8)
        while t <= newest:
            m = (t - MJD0).total_seconds() / 86400
            lines.append(f"{m:.6f} 0 {t.year} {t.month} {t.day} {t.hour} {t.minute} 0 0.500\n")
            t += timedelta(minutes=30)
        path.write_text("".join(lines))
        if mtime is not None:
            os.utime(path, (mtime, mtime))
        return path

    def write_plot(self, age_h: float) -> Path:
        path = self.files / "7_day_plot.png"
        path.write_bytes(b"png")
        t = time.time() - age_h * 3600
        os.utime(path, (t, t))
        return path


class TestPublicPlot(HealthTest):

    def test_fresh_plot_is_ok(self) -> None:
        self.write_spline(datetime(2026, 10, 19), mtime=time.time() - 9 * 3600)
        self.write_plot(age_h=9)
        sh.check_public_plot(STA)
        self.assertEqual([s for s, _ in self.findings("plot")], ["OK"])

    def test_stale_plot_warns(self) -> None:
        self.write_spline(datetime(2026, 10, 19), mtime=time.time() - 3 * 86400)
        self.write_plot(age_h=3 * 24)
        sh.check_public_plot(STA)
        (status, message), = self.findings("plot")
        self.assertEqual(status, "WARN")
        self.assertIn("not updating", message)

    def test_plot_older_than_the_record_warns(self) -> None:
        # The fit ran tonight, the plot did not: within 36 h, but wrong.
        self.write_spline(datetime(2026, 10, 19), mtime=time.time() - 9 * 3600)
        self.write_plot(age_h=33)
        sh.check_public_plot(STA)
        (status, message), = self.findings("plot")
        self.assertEqual(status, "WARN")
        self.assertIn("last night's plot failed", message)

    def test_missing_plot_warns(self) -> None:
        sh.check_public_plot(STA)
        self.assertEqual([s for s, _ in self.findings("plot")], ["WARN"])


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class TestTideCoverage(HealthTest):

    def write_tide(self, first: datetime, last: datetime) -> Path:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["time", "tide_m"])
        t = first
        while t <= last:
            ws.append([t, 0.5])
            t += timedelta(minutes=10)
        path = self.proj / "tides.xlsx"
        wb.save(path)
        return path

    def cfg(self) -> dict:
        return {"tide_model_file": str(self.proj / "tides.xlsx"),
                "tide_model_value_column": "tide_m"}

    def test_covering_prediction_is_ok(self) -> None:
        self.write_spline(datetime(2026, 10, 19, 23, 30))
        self.write_tide(datetime(2026, 10, 1), datetime(2026, 10, 31))
        sh.check_tide_coverage(STA, self.cfg())
        self.assertEqual([s for s, _ in self.findings("tide")], ["OK"])

    def test_prediction_ending_within_two_days_warns(self) -> None:
        self.write_spline(datetime(2026, 10, 19, 23, 30))
        self.write_tide(datetime(2026, 10, 1), datetime(2026, 10, 21))
        sh.check_tide_coverage(STA, self.cfg())
        (status, message), = self.findings("tide")
        self.assertEqual(status, "WARN")
        self.assertIn("within 2 days", message)

    def test_prediction_already_ended_warns(self) -> None:
        self.write_spline(datetime(2026, 10, 19, 23, 30))
        self.write_tide(datetime(2026, 10, 1), datetime(2026, 10, 15))
        sh.check_tide_coverage(STA, self.cfg())
        (status, message), = self.findings("tide")
        self.assertEqual(status, "WARN")
        self.assertIn("already missing", message)

    def test_prediction_starting_inside_the_window_warns(self) -> None:
        self.write_spline(datetime(2026, 10, 19, 23, 30))
        self.write_tide(datetime(2026, 10, 16), datetime(2026, 10, 31))
        sh.check_tide_coverage(STA, self.cfg())
        self.assertEqual([s for s, _ in self.findings("tide")], ["WARN"])

    def test_missing_file_warns(self) -> None:
        sh.check_tide_coverage(STA, self.cfg())
        self.assertEqual([s for s, _ in self.findings("tide")], ["WARN"])

    def test_no_tide_model_configured_is_not_a_warning(self) -> None:
        sh.check_tide_coverage(STA, {})
        self.assertEqual([s for s, _ in self.findings("tide")], ["OK"])

    def test_reader_without_openpyxl_agrees(self) -> None:
        # cron runs station_health.py under the system python, which may
        # not have openpyxl: the xlsx is then read directly.
        path = self.write_tide(datetime(2026, 10, 1, 0, 0), datetime(2026, 10, 31, 12, 0))
        expected = sh._tide_times(path, "time", "tide_m")
        got = sh._tide_times_stdlib(path, "time", "tide_m")
        self.assertEqual((expected[0], expected[-1]),
                         (datetime(2026, 10, 1), datetime(2026, 10, 31, 12)))
        self.assertEqual(got, expected)

    def test_times_without_a_tide_value_do_not_count(self) -> None:
        # plot_7day.load_tide draws a row only if its value is a number:
        # a sheet whose times run on past its values ends where the
        # values end.
        path = self.write_tide(datetime(2026, 10, 1), datetime(2026, 10, 31))
        wb = openpyxl.load_workbook(path)
        for row in wb.active.iter_rows(min_row=2):
            if row[0].value >= datetime(2026, 10, 15):
                row[1].value = None
        wb.save(path)
        self.write_spline(datetime(2026, 10, 19, 23, 30))
        sh.check_tide_coverage(STA, self.cfg())
        (status, message), = self.findings("tide")
        self.assertEqual(status, "WARN")
        self.assertIn("ends 2026-10-14 23:50", message)

    def test_renamed_value_column_warns(self) -> None:
        # The plot finds no tide_m column and draws no tide.
        self.write_spline(datetime(2026, 10, 19, 23, 30))
        self.write_tide(datetime(2026, 10, 1), datetime(2026, 10, 31))
        cfg = dict(self.cfg(), tide_model_value_column="EOT20_heightm")
        sh.check_tide_coverage(STA, cfg)
        (status, message), = self.findings("tide")
        self.assertEqual(status, "WARN")
        self.assertIn("no 'EOT20_heightm' column", message)

    def test_value_column_not_set_warns(self) -> None:
        self.write_tide(datetime(2026, 10, 1), datetime(2026, 10, 31))
        sh.check_tide_coverage(STA, {"tide_model_file": str(self.proj / "tides.xlsx")})
        (status, message), = self.findings("tide")
        self.assertEqual(status, "WARN")
        self.assertIn("tide_model_value_column", message)

    def test_hole_inside_the_window_warns(self) -> None:
        # Two predictions joined (merge_tide_files.py) span the window
        # but leave part of it without a tide.
        self.write_spline(datetime(2026, 10, 19, 23, 30))
        path = self.write_tide(datetime(2026, 10, 1), datetime(2026, 10, 31))
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        for row in list(ws.iter_rows(min_row=2)):
            if datetime(2026, 10, 16, 6) < row[0].value < datetime(2026, 10, 16, 18):
                row[1].value = None
        wb.save(path)
        sh.check_tide_coverage(STA, self.cfg())
        (status, message), = self.findings("tide")
        self.assertEqual(status, "WARN")
        self.assertIn("no predicted tide from 2026-10-16 06:00 to 2026-10-16 18:00", message)

    def test_stated_dimension_is_not_trusted(self) -> None:
        # openpyxl's read_only mode reads only as far as the sheet's
        # <dimension> says, and some writers state A1:A1; load_tide
        # (not read-only) reads every row.
        path = self.write_tide(datetime(2026, 10, 1), datetime(2026, 10, 31))
        with zipfile.ZipFile(path) as z:
            parts = {n: z.read(n) for n in z.namelist()}
        sheet = parts["xl/worksheets/sheet1.xml"].decode()
        start = sheet.index("<dimension ref=")
        end = sheet.index("/>", start) + 2
        parts["xl/worksheets/sheet1.xml"] = (
            sheet[:start] + '<dimension ref="A1:B2"/>' + sheet[end:]).encode()
        with zipfile.ZipFile(path, "w") as z:
            for name, data in parts.items():
                z.writestr(name, data)
        times = sh._tide_times(path, "time", "tide_m")
        self.assertEqual((times[0], times[-1]), (datetime(2026, 10, 1), datetime(2026, 10, 31)))


# A minimal xlsx written by hand, for the layouts openpyxl does not write
# but other spreadsheet programs do. Cells are given as XML.
_CT = ('<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/'
       'package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.'
       'openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="'
       'application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.'
       'openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/'
       'worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.'
       'spreadsheetml.worksheet+xml"/><Override PartName="/xl/styles.xml" ContentType="'
       'application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>{shared}'
       '</Types>')
_CT_SHARED = ('<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.'
              'openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>')
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_STYLES = (f'<styleSheet xmlns="{_MAIN}"><numFmts count="1"><numFmt numFmtId="164" '
           'formatCode="yyyy\\-mm\\-dd\\ hh:mm"/></numFmts><fonts count="1"><font><sz val="11"/>'
           '</font></fonts><fills count="2"><fill><patternFill patternType="none"/></fill><fill>'
           '<patternFill patternType="gray125"/></fill></fills><borders count="1"><border><left/>'
           '<right/><top/><bottom/><diagonal/></border></borders><cellStyleXfs count="1"><xf '
           'numFmtId="0"/></cellStyleXfs><cellXfs count="3"><xf numFmtId="0"/><xf numFmtId="164" '
           'applyNumberFormat="1"/><xf numFmtId="22" applyNumberFormat="1"/></cellXfs></styleSheet>')


def write_xlsx(path: Path, rows: list, date1904: bool = False, shared: list | None = None) -> Path:
    rels = [f'<Relationship Id="rId1" Type="{_REL}/worksheet" Target="worksheets/sheet1.xml"/>',
            f'<Relationship Id="rId2" Type="{_REL}/styles" Target="styles.xml"/>']
    if shared is not None:
        rels.append(f'<Relationship Id="rId3" Type="{_REL}/sharedStrings" '
                    f'Target="sharedStrings.xml"/>')
    pkg = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", _CT.format(shared=_CT_SHARED if shared is not None else ""))
        z.writestr("_rels/.rels", f'<Relationships xmlns="{pkg}"><Relationship Id="rId1" '
                                  f'Type="{_REL}/officeDocument" Target="xl/workbook.xml"/>'
                                  f'</Relationships>')
        z.writestr("xl/workbook.xml", f'<workbook xmlns="{_MAIN}" xmlns:r="{_REL}">'
                                      + ('<workbookPr date1904="1"/>' if date1904 else "")
                                      + '<sheets><sheet name="tide" sheetId="1" r:id="rId1"/>'
                                        '</sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", f'<Relationships xmlns="{pkg}">'
                                                 + "".join(rels) + "</Relationships>")
        z.writestr("xl/styles.xml", _STYLES)
        if shared is not None:
            z.writestr("xl/sharedStrings.xml", f'<sst xmlns="{_MAIN}">' + "".join(shared) + "</sst>")
        z.writestr("xl/worksheets/sheet1.xml", f'<worksheet xmlns="{_MAIN}"><sheetData>'
                                               + "".join(rows) + "</sheetData></worksheet>")
    return path


def header(time_col: str = "time", value_col: str = "tide_m") -> str:
    return (f'<row r="1"><c r="A1" t="inlineStr"><is><t>{time_col}</t></is></c>'
            f'<c r="B1" t="inlineStr"><is><t>{value_col}</t></is></c></row>')


def row(r: int, serial: float, value: float = 0.5, style: int = 1) -> str:
    return f'<row r="{r}"><c r="A{r}" s="{style}"><v>{serial}</v></c><c r="B{r}"><v>{value}</v></c></row>'


# 46314 is 2026-10-19 in Excel's usual 1900 date system; 1462 days less in 1904's
OCT19 = 46314


class TestTideReaderVariants(unittest.TestCase):
    """The xlsx reader for a python without openpyxl reads what openpyxl
    (and so plot_7day.load_tide) reads, on layouts other programs write."""

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def read(self, rows: list, **kw) -> list:
        path = write_xlsx(self.dir / "t.xlsx", rows, **kw)
        got = sh._tide_times_stdlib(path, "time", "tide_m")
        if openpyxl is not None:
            self.assertEqual(got, sh._tide_times(path, "time", "tide_m"))
        return got

    def test_1904_date_system(self) -> None:
        # Older Mac Excel counts from 1904-01-01; read from 1899-12-30
        # every time came out four years and a day early.
        got = self.read([header(), row(2, OCT19 - 1462), row(3, OCT19 - 1462 + 1.5)],
                        date1904=True)
        self.assertEqual(got, [datetime(2026, 10, 19), datetime(2026, 10, 20, 12)])

    def test_a_number_is_a_date_only_with_a_date_format(self) -> None:
        # The same serial in a cell with no date format is a plain number
        # to openpyxl, and the plot draws no tide at it.
        got = self.read([header(), row(2, OCT19, style=0), row(3, OCT19 + 0.25, style=2)])
        self.assertEqual(got, [datetime(2026, 10, 19, 6)])

    def test_cells_without_a_reference(self) -> None:
        # The r attribute is optional; a cell without one takes the next
        # column, a row without one the next row.
        rows = ['<row><c t="inlineStr"><is><t>time</t></is></c><c t="inlineStr"><is><t>tide_m'
                '</t></is></c></row>',
                f'<row><c s="1"><v>{OCT19}</v></c><c><v>0.5</v></c></row>',
                f'<row><c s="1"><v>{OCT19 + 1}</v></c><c><v>0.7</v></c></row>']
        self.assertEqual(self.read(rows), [datetime(2026, 10, 19), datetime(2026, 10, 20)])

    def test_shared_strings_with_rich_text_and_phonetic_runs(self) -> None:
        shared = ['<si><r><t>ti</t></r><r><t>me</t></r><rPh sb="0" eb="1"><t>X</t></rPh></si>',
                  '<si><t>tide_m</t></si>']
        rows = ['<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c></row>',
                row(2, OCT19)]
        self.assertEqual(self.read(rows, shared=shared), [datetime(2026, 10, 19)])

    def test_iso_date_cells(self) -> None:
        rows = [header(),
                '<row r="2"><c r="A2" t="d"><v>2026-10-19T00:00:00</v></c><c r="B2"><v>0.5</v></c></row>',
                '<row r="3"><c r="A3" t="d"><v>2026-10-20</v></c><c r="B3"><v>0.5</v></c></row>']
        self.assertEqual(self.read(rows), [datetime(2026, 10, 19)])

    def test_values_that_are_not_numbers_do_not_count(self) -> None:
        rows = [header(), row(2, OCT19),
                f'<row r="3"><c r="A3" s="1"><v>{OCT19 + 1}</v></c><c r="B3" t="str"><v>n/a</v></c></row>',
                f'<row r="4"><c r="A4" s="1"><v>{OCT19 + 2}</v></c><c r="B4" t="e"><v>#N/A</v></c></row>',
                f'<row r="5"><c r="A5" s="1"><v>{OCT19 + 3}</v></c><c r="B5"><f>B2</f></c></row>']
        self.assertEqual(self.read(rows), [datetime(2026, 10, 19)])

    def test_header_must_be_row_one(self) -> None:
        # load_tide reads the header from ws[1]: a header in row 2 gives
        # no tide on the plot, so it must not read as covered.
        rows = [header().replace('r="1"', 'r="2"').replace('A1', 'A2').replace('B1', 'B2'),
                row(3, OCT19)]
        with self.assertRaises(ValueError):
            self.read(rows)


class TestTideCoverageWithoutOpenpyxl(HealthTest):
    """The whole check as cron's system python runs it, with no openpyxl."""

    def test_1904_workbook_covering_the_window_is_ok(self) -> None:
        self.write_spline(datetime(2026, 10, 19, 23, 30))
        rows = [header()] + [row(k + 2, OCT19 - 1462 - 20 + k / 24) for k in range(24 * 32)]
        path = write_xlsx(self.proj / "tides.xlsx", rows, date1904=True)
        with unittest.mock.patch.dict(sys.modules, {"openpyxl": None}):
            sh.check_tide_coverage(STA, {"tide_model_file": str(path),
                                         "tide_model_value_column": "tide_m"})
        (status, message), = self.findings("tide")
        self.assertEqual(status, "OK", message)
        self.assertIn("covers 2026-09-29 to 2026-10-30", message)


class TestArcStats(HealthTest):
    """Storm amplitude damping: arcs drop below gnssir's limits."""

    def write_results(self, day: datetime, n: int, amp: float, pknoise: float = 4.0) -> None:
        doy = day.timetuple().tm_yday
        path = (self.proj / "products" / "refl_code" / str(day.year) / "results"
                / STA / f"{doy:03d}.txt")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            "% year doy RH sat UTCtime Azim Amp eminO emaxO NumbOf freq rise EdotF "
            "PkNoise DelT MJD refr-appl\n"
            + "".join(f"{day.year} {doy} 19.000 {i % 32 + 1} {24 * i / n:.3f} 80.00 "
                      f"{amp + (i % 3) - 1:.2f} 5.05 14.95 200 1 1 0.7 {pknoise:.2f} 30.0 "
                      f"61000.0 1\n" for i in range(n)))

    def week(self, last_n: int, last_amp: float) -> None:
        for k in range(8):
            day = datetime(2026, 10, 12) + timedelta(days=k)
            if k < 7:
                self.write_results(day, n=400, amp=12.0)
            else:
                self.write_results(day, n=last_n, amp=last_amp, pknoise=3.1)
        gr.write_status(self.files, ok=True)

    def test_status_holds_per_day_arc_statistics(self) -> None:
        self.week(last_n=390, last_amp=11.0)
        st = json.loads((self.files / "fit" / gr.STATUS_NAME).read_text())
        self.assertEqual(len(st["arc_stats"]), 8)
        self.assertEqual(st["arc_stats"][0],
                         {"day": "2026-10-12", "arcs": 400, "median_amp": 12.0,
                          "median_pknoise": 4.0})
        self.assertEqual(st["arc_stats"][-1]["day"], "2026-10-19")
        self.assertEqual(st["arc_stats"][-1]["median_pknoise"], 3.1)

    def test_steady_day_is_ok(self) -> None:
        self.week(last_n=390, last_amp=11.0)
        sh.check_arc_stats(STA)
        self.assertEqual([s for s, _ in self.findings("arcs")], ["OK"])

    def test_arc_count_drop_warns(self) -> None:
        self.week(last_n=150, last_amp=11.0)
        sh.check_arc_stats(STA)
        (status, message), = self.findings("arcs")
        self.assertEqual(status, "WARN")
        self.assertIn("drop in arcs", message)

    def test_amplitude_drop_warns(self) -> None:
        self.week(last_n=390, last_amp=7.0)
        sh.check_arc_stats(STA)
        (status, message), = self.findings("arcs")
        self.assertEqual(status, "WARN")
        self.assertIn("amplitude", message)

    def test_too_few_days_says_nothing(self) -> None:
        self.write_results(datetime(2026, 10, 19), n=10, amp=5.0)
        gr.write_status(self.files, ok=True)
        sh.check_arc_stats(STA)
        self.assertEqual(self.findings("arcs"), [])

    def test_no_status_file_says_nothing(self) -> None:
        sh.check_arc_stats(STA)
        self.assertEqual(self.findings("arcs"), [])

    def results_dir(self, day: datetime) -> Path:
        path = self.proj / "products" / "refl_code" / str(day.year) / "results" / STA
        path.mkdir(parents=True, exist_ok=True)
        return path

    def kept_nothing(self, day: datetime) -> None:
        """What gnssrefl 4.1.5 leaves for a day it kept no arc of: no
        <doy>.txt (retrieve_rh deletes it), the rejected arcs in failQC/."""
        failqc = self.results_dir(day) / "failQC"
        failqc.mkdir(exist_ok=True)
        (failqc / f"{day.timetuple().tm_yday:03d}.txt").write_text(
            "% year doy RH sat UTCtime Azim Amp ...\n"
            "2026 292 nan 5 3.100 80.00 nan 5.05 14.95 200 1 1 0.7 nan 30.0 61000.0 1\n")

    def test_a_day_with_no_arc_kept_warns(self) -> None:
        # The worst a storm can do has no results file: it used to be left
        # out, so the day before it was compared and read OK.
        for k in range(7):
            self.write_results(datetime(2026, 10, 12) + timedelta(days=k), n=400, amp=12.0)
        self.kept_nothing(datetime(2026, 10, 19))
        gr.write_status(self.files, ok=True)
        st = json.loads((self.files / "fit" / gr.STATUS_NAME).read_text())
        self.assertEqual(st["arc_stats"][-1], {"day": "2026-10-19", "arcs": 0,
                                               "median_amp": None, "median_pknoise": None})
        sh.check_arc_stats(STA)
        (status, message), = self.findings("arcs")
        self.assertEqual(status, "WARN")
        self.assertIn("2026-10-19: no arcs kept", message)

    def test_a_no_data_marker_is_a_day_with_no_arc_kept(self) -> None:
        # recover_missing_days.sh's marker (unpadded day of year); inside
        # the baseline it counts for the arcs, not for the amplitude
        for k in range(8):
            day = datetime(2026, 10, 12) + timedelta(days=k)
            if k == 4:
                (self.results_dir(day) / f"{day.timetuple().tm_yday}.no_data").write_text(
                    "2026-10-20T03:00:00Z -- no usable retrievals\n")
            else:
                self.write_results(day, n=400 if k < 7 else 390, amp=12.0 if k < 7 else 11.0)
        gr.write_status(self.files, ok=True)
        st = json.loads((self.files / "fit" / gr.STATUS_NAME).read_text())
        self.assertEqual([d["arcs"] for d in st["arc_stats"]], [400] * 4 + [0] + [400] * 2 + [390])
        sh.check_arc_stats(STA)
        (status, message), = self.findings("arcs")
        self.assertEqual(status, "OK", message)
        self.assertIn("previous 7 days: 400 arcs, Amp 12.0", message)

    def test_a_day_gnssir_never_ran_for_is_not_counted(self) -> None:
        # no SNR file yet (the orbit not out): nothing in results/ at all
        self.week(last_n=390, last_amp=11.0)
        self.results_dir(datetime(2026, 10, 20))
        gr.write_status(self.files, ok=True)
        st = json.loads((self.files / "fit" / gr.STATUS_NAME).read_text())
        self.assertEqual(st["arc_stats"][-1]["day"], "2026-10-19")


class TestRecordStatusAge(HealthTest):
    """record_status.json is what the last update that wrote it said: an
    update that did not run leaves the night before's ok=True."""

    def setUp(self) -> None:
        super().setUp()
        for k in range(8):
            day = datetime(2026, 10, 12) + timedelta(days=k)
            doy = day.timetuple().tm_yday
            path = (self.proj / "products" / "refl_code" / str(day.year) / "results"
                    / STA / f"{doy:03d}.txt")
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(f"% results\n{day.year} {doy} 19.000 5 12.000 80.00 12.00\n")
        self.write_spline(datetime(2026, 10, 19, 23, 30))

    def status(self, age_h: float) -> None:
        written = sh.utcnow() - timedelta(hours=age_h)
        (self.files / "fit").mkdir(exist_ok=True)
        (self.files / "fit" / gr.STATUS_NAME).write_text(json.dumps(
            {"ok": True, "time": written.strftime("%Y-%m-%dT%H:%M:%SZ")}))

    def test_last_nights_status_is_ok(self) -> None:
        self.status(age_h=9)
        sh.check_water_level_record(STA)
        self.assertEqual([s for s, _ in self.findings("record")], ["OK"])

    def test_status_not_written_for_days_warns(self) -> None:
        self.status(age_h=3 * 24)
        sh.check_water_level_record(STA)
        (status, message), = self.findings("record")
        self.assertEqual(status, "WARN")
        self.assertIn("3.0 days ago", message)


if __name__ == "__main__":
    unittest.main()
