#!/usr/bin/env python3
"""
make_tide_prediction.py

Makes the predicted-tide spreadsheet that the public 7-day plot draws,
so the plot no longer depends on a hand-made file that runs out.
OPT-IN: nothing runs this unless you ask for it (see SWITCHING IT ON).

WHY

The orange "Predicted tide" line on 7_day_plot.png, and every
departure marker on it, come from tide_model_file in station.json.
That is a hand-made EOT20 spreadsheet (marconi_tides_*.xlsx) covering
a fixed period. When it ends, both disappear from the plot:
plot_7day.load_tide finds nothing, and station_health warns two days
before. This script makes the same file automatically for a rolling
30 days. By default the window runs from 10 days before today to 20
days ahead. That covers the plot's week and the health check's two
days of look-ahead, and leaves a usable file even after a few weeks
of failed runs.

THE SAME FILE, SO NOTHING THAT READS IT CHANGES

The output is an .xlsx whose first sheet has a header row, a time
column and one height column, with one row every 10 minutes. That is
the layout of the spreadsheet it replaces, and what plot_7day.load_tide,
station_health and the comparison tools read. The column names default
to what station.json already uses: tide_model_time_column for the time
("time"). The height column is "EOT20_heightm" for EOT20, as in the
spreadsheet, and "NOAA<station>_heightm" for NOAA. A second sheet,
"about", says where the numbers came from. Every reader takes only
the first sheet.

Heights are metres about the model's own mean sea level (a zero-mean
tide), as the EOT20 spreadsheet's are. tide_model_navd88_offset
therefore keeps its meaning. Each value is the instantaneous height
at its time, not an average over the 10 minutes.

TIME IS UTC, EXPLICITLY

The spreadsheet's times are naive and only assumed to be UTC, and its
two source files disagree by one 10-minute row (merge_tide_files.py).
Here the times are made in UTC and handed to the source as UTC: to
pyTMD with standard "UTC" (it adds the leap seconds itself), and to
NOAA with time_zone=gmt. They are written naive, because every reader
expects naive times, and the "about" sheet says they are UTC. All
fall on whole 10-minute marks.

SOURCES (--source, or "source" in station.json's "tide_prediction")

  pytmd   EOT20 (Hart-Davis et al. 2021, Earth Syst. Sci. Data 13,
          3869-3884, doi:10.5194/essd-13-3869-2021), the model the
          spreadsheet was made from, computed with pyTMD (Sutterley
          et al., doi:10.5281/zenodo.5555395) at the station's latitude
          and longitude. This is the ocean tide, not the load tide:
          the antenna moves with the loaded crust, as a tide gauge
          does. It needs EOT20's ocean_tides files (doi:10.17882/79489)
          in <directory>/EOT20/ocean_tides/ and pyTMD in the virtual
          environment. Tested: pyTMD 3.0.9 (Python 3.9 or later, with
          dask) and 2.2.9.1. gnssrefl_venv is pinned (requirements.txt)
          and pyTMD 3 brings in xarray, pandas and scipy, so install it
          against the current versions, which makes pip refuse rather
          than upgrade anything:
              pip freeze > before.txt
              pip install 'pyTMD==3.0.9' dask -c before.txt
          Memory: pyTMD 3 loads every constituent's global grid, about
          2.3-2.5 GB at EOT20's 1/8 degree whether or not it crops;
          pyTMD 2 about 0.5 GB (both measured on a synthetic grid of
          that size, not on the NUC).
          The station is on a narrow strip of land between the
          Atlantic and Cape Cod Bay, which have different tides. If
          the model treats the station's grid cell as land, pyTMD
          fills it from the nearest ocean cell within --extrapolate-km
          (the run says so, and the "about" sheet records it). When
          that cell is on the bay side, the prediction is the bay's
          tide, not the open coast's. Even with ocean all round, the
          linear interpolation between grid nodes can mix a bay node
          with an Atlantic one. --validate's height ratio and RMS show
          either. The fix is a "point" ([lat, lon]) a little offshore
          on the Atlantic side.
  noaa    NOAA CO-OPS harmonic predictions (Parker 2007, Tidal Analysis
          and Prediction, NOAA Special Publication NOS CO-OPS 3) for a
          harmonic station: datum MSL, metric, time_zone=gmt. The
          default station is 8447435, Chatham Lydia Cove, the gauge
          water_level_qc.py already uses. Lydia Cove is a harbour
          behind a barrier beach, not the open coast here, so its
          range and timing differ; --validate measures by how much.
          --noaa-time-offset-minutes and --noaa-height-ratio apply the
          two corrections NOAA itself uses for its subordinate
          stations: a time shift and a height ratio. Both are off by
          default. NOAA's MSL is the gauge's mean over its tidal datum
          epoch, not today's. That does not move the curve, which has a
          zero mean like EOT20's, but its seasonal terms and the
          harbour can put it a few cm off the EOT20 one.
          tide_model_navd88_offset was measured for the EOT20
          spreadsheet: before tide_model_file points at a NOAA file,
          check --validate's mean difference and adjust the offset by
          it if it is not near zero.
  auto    pytmd if it works on this machine, otherwise noaa. It says
          which it used and why. Once tide_model_file is this file, it
          uses the source whose height column tide_model_value_column
          names, and otherwise, once it has written the file, it keeps
          that file's source: switching on a night pyTMD fails would
          rename the height column, and the plot would lose its
          predicted tide. Delete the file, or give --source, to choose
          again. An explicit "source" is safer than auto.

VALIDATION (--validate)

Compares a prediction with the existing spreadsheet(s) over their
overlap and writes nothing. It reports:
  - the mean difference, the RMS (raw and with the mean removed), the
    largest difference and the correlation;
  - the time lag between -6 and +6 h that fits best, and the RMS at
    that lag;
  - the height ratio at that lag (spreadsheet range / prediction
    range).
How to read the lag, against pytmd/EOT20 (the same model, so the lag
is absolute):
  - within 2.5 min of zero: the two agree in time;
  - about 10 min: one 10-minute row (merge_tide_files.py's ambiguity);
  - about -4 or -5 h: Eastern time written as UTC (+4 or +5 h: the
    times moved the other way, e.g. converted to UTC twice);
  - anything else: a real difference between the two predictions
    (another model or place), not a clock error.
Only then does the report say how many minutes to add to a
spreadsheet. Against noaa the lag includes the harbour's real delay,
so it is not read as a clock error at all. Give both spreadsheets
(--reference OLD NEW) and it also prints the difference between their
lags, which holds for either source: that is the shift
merge_tide_files.py --shift-old-minutes needs.

SAFETY

  - The file is written under a temporary name and then renamed into
    place, so a failed run leaves the previous file untouched.
  - It never overwrites a spreadsheet it did not make, that is, one
    without its "about" sheet, such as the hand-made file. --force
    overrides this.
  - It keeps history by default ("keep_history", --no-keep-history to
    stop): the earlier rows of its own previous file stay when that
    file came from the same source, and the new prediction wins where
    the two overlap. The file grows by about 52,000 rows a year.
    Everything else that reads tide_model_file reads the WHOLE file:
    Step 6 of process_and_plot.sh, the morning mail's statistics and
    the tide column of the monthly CSVs that maintenance/
    s3UploadTimeseries.sh publishes. That script uploads each month
    again for 10 days after it ends, and every month of a year again
    when gnss_record.py freezes the year; a month outside the file
    gets an empty tide column. So without history, each published
    month would keep only its last day or so of tide, and the months
    before the switch none at all.
  - For the same reason its first run (no file of its own yet) copies,
    as they are, the rows of the spreadsheet tide_model_file names
    that come before its window, when that spreadsheet has the same
    height column, and predicts from where the spreadsheet ends. The
    "about" sheet records it ("seeded_from"). Months already published
    then keep their tide values when they are uploaded again.

SWITCHING IT ON (on the NUC)

  1. Validate it against the spreadsheets you have (see VALIDATION).
  2. Refresh it nightly by adding to station.json, for example:
         "tide_prediction": {"source": "pytmd",
                             "pytmd_directory": "/home/argus_user/tide_models"}
     Optional keys: "enabled" (false pauses it), "file", "model",
     "point", "extrapolate_km", "noaa_station",
     "noaa_time_offset_minutes", "noaa_height_ratio", "keep_history"
     and "value_column" (STATION_JSON_REFERENCE.md).
     process_and_plot.sh then refreshes the file just before it draws
     the 7-day plot. Without a "tide_prediction" section nothing runs.
     Let it run once while tide_model_file is still the spreadsheet,
     so the file starts with the spreadsheet's rows (SAFETY).
  3. Point the plot at it with "tide_model_file" (the FULL path this
     script prints: the plot looks a relative path up from wherever
     it was started) and "tide_model_value_column" ("EOT20_heightm").
     Until you do, nothing the public sees changes. From then on the
     file is also the tide column of the monthly CSVs on S3, so leave
     keep_history on. To go back, point tide_model_file at the
     spreadsheet again.

Usage:
    python3 make_tide_prediction.py --source pytmd --pytmd-directory ~/tide_models
    python3 make_tide_prediction.py --source noaa --output tides/tide_prediction.xlsx
    python3 make_tide_prediction.py --validate --source pytmd \\
        --pytmd-directory ~/tide_models \\
        --reference marconi_tides_sherwood_20260715-20260831.xlsx marconi_tides_sherwood.xlsx
"""

from __future__ import annotations

import argparse
import json
import math
import os
import sys
import urllib.request
from datetime import datetime, timedelta, timezone
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parent))
from station_datum import PROJECT_DIR, load_station  # noqa: E402

STEP_MINUTES = 10                     # the spreadsheet's row spacing
DAYS = 30                             # length of the rolling window
DAYS_BACK = 10                        # of which before today (UTC)
DEFAULT_MODEL = "EOT20"
DEFAULT_NOAA_STATION = "8447435"      # Chatham, Lydia Cove, MA (water_level_qc.GAUGE_STATION)
EXTRAPOLATE_KM = 10.0                 # pyTMD's own default cutoff
DEFAULT_OUTPUT = PROJECT_DIR / "tides" / "tide_prediction.xlsx"
GENERATOR = "make_tide_prediction.py"
ABOUT_SHEET = "about"
UNIX_EPOCH = datetime(1970, 1, 1)

NOAA_API = ("https://api.tidesandcurrents.noaa.gov/api/prod/datagetter?product=predictions"
            "&station={station}&begin_date={begin}&end_date={end}&datum=MSL&units=metric"
            "&time_zone=gmt&interval={interval}&format=json&application=gnssir_station_tide")
NOAA_CHUNK_DAYS = 30                  # the API serves a limited span per request

# Validation
MAX_LAG_MINUTES = 360                 # +/-6 h: catches a 4-5 h time-zone error
MIN_OVERLAP_ROWS = 144                # one day of 10-minute rows
LAG_SAME_MIN = 2.5                    # within this, the two agree in time
TIME_ZONE_HOURS = (4, 5)              # EDT and EST behind UTC
TIME_ZONE_TOL_MIN = 15


def utcnow() -> datetime:
    """Now as naive UTC. A function so the tests can fix the clock."""
    return datetime.now(timezone.utc).replace(tzinfo=None)


def rolling_window(now: datetime, days_back: int = DAYS_BACK, days: int = DAYS):
    """(start, end) of the rolling prediction, naive UTC. It starts at
    midnight UTC days_back days before `now` and runs for `days`
    days; end itself is not included."""
    start = datetime(now.year, now.month, now.day) - timedelta(days=days_back)
    return start, start + timedelta(days=days)


def grid_times(start: datetime, end: datetime, step_minutes: int = STEP_MINUTES):
    """Every whole step_minutes mark from start (moved forward onto
    the grid if it falls between marks) up to, but not including, end."""
    step = timedelta(minutes=step_minutes)
    midnight = datetime(start.year, start.month, start.day)
    t = midnight + step * math.ceil((start - midnight) / step)
    times = []
    while t < end:
        times.append(t)
        t += step
    return times


def _seconds(times) -> np.ndarray:
    return np.array([(t - UNIX_EPOCH).total_seconds() for t in times], dtype=float)


def _interp_no_gaps(xq, x, y, max_gap_s: float) -> np.ndarray:
    """y(x) linearly interpolated to xq, NaN outside x's range and
    inside any gap between samples wider than max_gap_s (a value that
    falls exactly on a sample is always kept). x must be sorted."""
    xq = np.asarray(xq, dtype=float)
    out = np.interp(xq, x, y, left=np.nan, right=np.nan)
    if len(x) < 2:
        return out
    i = np.searchsorted(x, xq, side="left")
    hi = np.clip(i, 0, len(x) - 1)
    lo = np.clip(i - 1, 0, len(x) - 1)
    on_sample = x[hi] == xq
    out[((x[hi] - x[lo]) > max_gap_s) & ~on_sample] = np.nan
    return out


# ---------------------------------------------------------------------
# SOURCES
# ---------------------------------------------------------------------

def pytmd_prediction(times, lat: float, lon: float, directory=None,
                     model: str = DEFAULT_MODEL, extrapolate_km: float = EXTRAPOLATE_KM,
                     infer_minor: bool = True, info: dict | None = None):
    """Ocean-tide heights (m) at the naive-UTC `times`, from a pyTMD
    model (EOT20 by default) at one point. If `info` is a dict,
    info["land_fill"] is set True when the point had to be filled from
    the ocean nearby.

    The times go to pyTMD as UTC seconds since 1970 with standard
    "UTC", so pyTMD itself applies the leap seconds. Minor
    constituents are inferred from the major ones, as pyTMD does by
    default (infer_minor).

    pyTMD 3 takes lower-case keywords and pyTMD 2 upper-case ones.
    Both accept any keyword without complaint, so the version is told
    from the function's signature rather than by trial.

    The point is first interpolated from the model grid as it is. Only
    if that gives no value (a grid cell next to it is land) is it
    filled from the nearest ocean within extrapolate_km. pyTMD 3 needs
    dask for that fill, and for cropping the global grids to 1 degree
    around the point (which does not measurably lower its peak memory),
    and does not install dask itself; without dask, a point whose grid
    cells are all ocean still works. Raises if pyTMD or the model files
    are missing, or if no ocean value lies within extrapolate_km of the
    point."""
    import importlib.util
    import inspect
    import pyTMD.compute                       # ImportError: not installed

    if lat is None or lon is None:
        raise ValueError("no latitude/longitude (station.json latitude, longitude "
                         "or tide_prediction \"point\")")
    n = len(times)
    secs = _seconds(times)
    x, y = np.full(n, float(lon)), np.full(n, float(lat))
    epoch = (1970, 1, 1, 0, 0, 0)
    params = inspect.signature(pyTMD.compute.tide_elevations).parameters
    have_dask = importlib.util.find_spec("dask") is not None

    def heights(extrapolate: bool) -> np.ndarray:
        if "directory" in params:              # pyTMD 3
            kw = dict(model=model, crs=4326, epoch=epoch, type="trajectory",
                      standard="UTC", method="linear", extrapolate=extrapolate,
                      cutoff=extrapolate_km, crop=have_dask, buffer=1.0,
                      infer_minor=infer_minor)
            if directory:
                kw["directory"] = str(directory)
        else:                                  # pyTMD 2
            # No cropping: pyTMD 2.2's crop fails on a 0-360 degree grid
            # such as EOT20's for a point west of Greenwich (it crops
            # twice). It reads one constituent at a time, so the whole
            # grid costs little memory.
            kw = dict(MODEL=model, EPSG=4326, EPOCH=epoch, TYPE="drift", TIME="UTC",
                      METHOD="linear", EXTRAPOLATE=extrapolate, CUTOFF=extrapolate_km,
                      CROP=False, INFER_MINOR=infer_minor)
            if directory:
                kw["DIRECTORY"] = str(directory)
        out = pyTMD.compute.tide_elevations(x, y, secs, **kw)
        if np.ma.isMaskedArray(out):
            values = np.ma.filled(out.astype(float), np.nan)
        else:
            values = np.asarray(getattr(out, "values", out), dtype=float)
        values = values.reshape(-1)
        if values.size != n:
            raise RuntimeError(f"pyTMD returned {values.size} heights for {n} times")
        return values

    values = heights(extrapolate=False)
    if not np.isfinite(values).all():
        if info is not None:
            info["land_fill"] = True
        try:
            values = heights(extrapolate=True)
        except ImportError as exc:
            raise RuntimeError(f"the point {lat:.4f}, {lon:.4f} is next to land in {model}, "
                               f"and filling it from the ocean nearby needs dask "
                               f"(pip install dask), or give a point a little offshore: "
                               f"{exc}") from exc
    if not np.isfinite(values).all():
        raise RuntimeError(f"{model} has no ocean value within {extrapolate_km:g} km of "
                           f"{lat:.4f}, {lon:.4f}: give a point a little offshore "
                           f"(tide_prediction \"point\", or --lat/--lon) or a larger "
                           f"--extrapolate-km")
    return values


def _pytmd_version() -> str:
    try:
        import pyTMD
        return str(getattr(pyTMD, "__version__", "") or getattr(pyTMD.version, "version", ""))
    except Exception:
        return ""


def _noaa_fetch(station: str, first_day: datetime, last_day: datetime, interval: str,
                urlopen) -> dict:
    """{naive UTC time: height} of CO-OPS predictions for the whole UTC
    days first_day..last_day, in requests of NOAA_CHUNK_DAYS days.

    Each answer is checked: it must start at its first day's midnight,
    end on its last day, and every time must fall on the interval's
    grid. A short or malformed answer raises instead of leaving a
    hole in the file."""
    rows = {}
    day = first_day
    while day <= last_day:
        stop = min(day + timedelta(days=NOAA_CHUNK_DAYS - 1), last_day)
        url = NOAA_API.format(station=station, begin=day.strftime("%Y%m%d"),
                              end=stop.strftime("%Y%m%d"), interval=interval)
        data = json.loads(urlopen(url, timeout=60).read().decode())
        if not isinstance(data, dict) or not data.get("predictions"):
            message = data.get("error", data) if isinstance(data, dict) else data
            raise RuntimeError(f"NOAA CO-OPS returned no predictions for {station} "
                               f"(interval {interval}): {str(message)[:200]}")
        chunk = {}
        for p in data["predictions"]:
            chunk[datetime.strptime(p["t"].strip(), "%Y-%m-%d %H:%M")] = float(p["v"])
        step = int(interval)
        if (min(chunk) != day or max(chunk).date() != stop.date()
                or any(t.minute % step or t.second for t in chunk)):
            raise RuntimeError(f"NOAA CO-OPS answer for {station} does not cover "
                               f"{day:%Y-%m-%d} to {stop:%Y-%m-%d} on a {step}-minute grid "
                               f"(got {min(chunk)} to {max(chunk)})")
        rows.update(chunk)
        day = stop + timedelta(days=1)
    return rows


def noaa_prediction(times, station: str = DEFAULT_NOAA_STATION,
                    time_offset_minutes: float = 0.0, height_ratio: float = 1.0,
                    urlopen=None):
    """Heights (m above the station's MSL) at the naive-UTC `times`
    from NOAA CO-OPS harmonic predictions, requested in UTC
    (time_zone=gmt).

    The 10-minute interval is asked for first, which gives the
    predictions at exactly these times. If the API refuses it, the
    6-minute default is used and interpolated, which is within 0.5 mm
    for this site's tide.

    time_offset_minutes and height_ratio are NOAA's subordinate-station
    corrections (Parker 2007): the site's tide is the station's,
    `time_offset_minutes` later and `height_ratio` times as large.
    Defaults are no change."""
    urlopen = urlopen or urllib.request.urlopen
    shift = timedelta(minutes=float(time_offset_minutes))
    want = [t - shift for t in times]
    first = datetime(min(want).year, min(want).month, min(want).day)
    # one extra day, so the last times always have a sample after them
    last = datetime(max(want).year, max(want).month, max(want).day) + timedelta(days=1)
    try:
        rows = _noaa_fetch(station, first, last, "10", urlopen)
    except Exception as exc:
        print(f"  NOAA 10-minute predictions not available ({exc}); "
              f"using the 6-minute ones, interpolated")
        rows = _noaa_fetch(station, first, last, "6", urlopen)
    rt = sorted(rows)
    values = _interp_no_gaps(_seconds(want), _seconds(rt), np.array([rows[t] for t in rt]),
                             max_gap_s=STEP_MINUTES * 60)
    if not np.isfinite(values).all():
        missing = [t for t, v in zip(times, values) if not np.isfinite(v)]
        raise RuntimeError(f"NOAA predictions for {station} are missing at {len(missing)} "
                           f"time(s), from {missing[0]}")
    return float(height_ratio) * values


# ---------------------------------------------------------------------
# OPTIONS AND THE PREDICTION ITSELF
# ---------------------------------------------------------------------

def resolve_options(args, cfg: dict) -> dict:
    """Each setting from the command line, else from station.json's
    "tide_prediction", else the default."""
    tp = cfg.get("tide_prediction")
    tp = tp if isinstance(tp, dict) else {}

    def pick(cli, key, default):
        if cli is not None:
            return cli
        return tp[key] if tp.get(key) is not None else default

    source = str(pick(args.source, "source", "auto")).lower()
    if source not in ("auto", "pytmd", "noaa"):
        raise SystemExit(f"tide_prediction source {source!r} is not auto, pytmd or noaa")
    point = tp.get("point")
    if point is not None and not (isinstance(point, (list, tuple)) and len(point) == 2):
        raise SystemExit(f"tide_prediction \"point\" must be [latitude, longitude], "
                         f"not {point!r}")
    keep_history = pick(args.keep_history, "keep_history", True)
    if not isinstance(keep_history, bool):
        raise SystemExit(f"tide_prediction \"keep_history\" must be true or false, "
                         f"not {keep_history!r}")
    lat = args.lat if args.lat is not None else (point[0] if point else cfg.get("latitude"))
    lon = args.lon if args.lon is not None else (point[1] if point else cfg.get("longitude"))
    output = Path(pick(args.output, "file", DEFAULT_OUTPUT)).expanduser()
    if not output.is_absolute():
        output = PROJECT_DIR / output
    return {
        "source": source,
        "output": output,
        "lat": float(lat) if lat is not None else None,
        "lon": float(lon) if lon is not None else None,
        "pytmd_directory": pick(args.pytmd_directory, "pytmd_directory", None),
        "model": str(pick(args.model, "model", DEFAULT_MODEL)),
        "extrapolate_km": float(pick(args.extrapolate_km, "extrapolate_km", EXTRAPOLATE_KM)),
        "noaa_station": str(pick(args.noaa_station, "noaa_station", DEFAULT_NOAA_STATION)),
        "noaa_time_offset_minutes": float(pick(args.noaa_time_offset_minutes,
                                               "noaa_time_offset_minutes", 0.0)),
        "noaa_height_ratio": float(pick(args.noaa_height_ratio, "noaa_height_ratio", 1.0)),
        "keep_history": keep_history,
        "value_column": pick(args.value_col, "value_column", None),
        "time_column": args.time_col or cfg.get("tide_model_time_column") or "time",
    }


def predict(times, opts: dict):
    """(heights, about) for `times` from the configured source. `about`
    describes the source for the "about" sheet; about["source_key"]
    names it exactly, so --keep-history never joins two sources."""
    source = opts["source"]
    if source in ("auto", "pytmd"):
        try:
            info = {}
            values = pytmd_prediction(times, opts["lat"], opts["lon"], opts["pytmd_directory"],
                                      opts["model"], opts["extrapolate_km"], info=info)
            name = f"pyTMD {_pytmd_version()}".strip()
            about = {
                "source": f"{name}, {opts['model']} ocean tide",
                "point": f"{opts['lat']:.5f}, {opts['lon']:.5f} (latitude, longitude)",
                "heights": "metres about the model's mean sea level (a zero-mean tide); "
                           "add tide_model_navd88_offset for NAVD88",
                "source_key": f"pytmd:{opts['model']}:{opts['lat']:.5f},{opts['lon']:.5f}",
                "value_column": f"{opts['model']}_heightm",
            }
            if info.get("land_fill"):
                about["land_fill"] = (f"the point's grid cell is land in {opts['model']}: "
                                      f"filled from the nearest ocean node within "
                                      f"{opts['extrapolate_km']:g} km, not interpolated")
                print(f"  filled from the nearest {opts['model']} ocean node within "
                      f"{opts['extrapolate_km']:g} km (the station cell is land in the model)")
            return values, about
        except Exception as exc:
            if source == "pytmd":
                raise
            print(f"  pyTMD/{opts['model']} not usable here ({type(exc).__name__}: {exc}); "
                  f"using NOAA CO-OPS predictions for {opts['noaa_station']}")
    station = opts["noaa_station"]
    offset, ratio = opts["noaa_time_offset_minutes"], opts["noaa_height_ratio"]
    values = noaa_prediction(times, station, offset, ratio)
    return values, {
        "source": f"NOAA CO-OPS harmonic predictions, station {station}",
        "point": f"NOAA station {station}",
        "heights": "metres above the station's mean sea level (datum MSL of NOAA's tidal "
                   "datum epoch, a zero-mean tide); tide_model_navd88_offset was measured "
                   "for EOT20: check it against --validate's mean difference",
        "adjustment": (f"{offset:+g} min, x{ratio:g} (subordinate-station correction)"
                       if offset or ratio != 1.0 else "none"),
        "source_key": f"noaa:{station}:{offset:g}:{ratio:g}",
        "value_column": f"NOAA{station}_heightm",
    }


# ---------------------------------------------------------------------
# THE FILE
# ---------------------------------------------------------------------

def write_prediction(path: Path, times, values, time_col: str, value_col: str,
                     about: dict) -> None:
    """The spreadsheet: sheet 1 is the header row plus one row per
    time; sheet 2 ("about") holds the provenance. Written under a
    temporary name and renamed, so readers never see half a file and a
    failure leaves the old one."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "prediction"
    ws.append([time_col, value_col])
    for t, v in zip(times, values):
        ws.append([t, round(float(v), 4)])
    info = wb.create_sheet(ABOUT_SHEET)
    info.append(["generator", GENERATOR])
    for key, value in about.items():
        info.append([key, value])
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.stem}.tmp{path.suffix}")
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if tmp.exists():
            tmp.unlink()


def read_prediction(path: Path, time_col: str = "time", value_col: str | None = None):
    """(times, values, about, value_col) from a tide spreadsheet's
    first sheet, read the way plot_7day.load_tide reads it: rows whose
    time is a datetime and whose height is a number. value_col
    defaults to the first column that is not the time. `about` is the
    "about" sheet as a dict, or {} for a file this script did not
    make."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb[wb.sheetnames[0]].iter_rows(values_only=True)
        header = list(next(rows, None) or [])
        if time_col not in header:
            raise ValueError(f"{path.name}: no {time_col!r} column (found {header})")
        if value_col is None:
            others = [h for h in header if h is not None and h != time_col]
            if not others:
                raise ValueError(f"{path.name}: no height column (found {header})")
            value_col = others[0]
        if value_col not in header:
            raise ValueError(f"{path.name}: no {value_col!r} column (found {header})")
        ti, vi = header.index(time_col), header.index(value_col)
        times, values = [], []
        for r in rows:
            if len(r) <= max(ti, vi):
                continue
            t, v = r[ti], r[vi]
            if isinstance(t, datetime) and isinstance(v, (int, float)) and math.isfinite(v):
                times.append(t)
                values.append(float(v))
        about = {}
        if ABOUT_SHEET in wb.sheetnames:
            for r in wb[ABOUT_SHEET].iter_rows(values_only=True):
                if r and r[0] is not None:
                    about[str(r[0])] = r[1] if len(r) > 1 else None
    finally:
        wb.close()
    return times, np.asarray(values, dtype=float), about, value_col


def is_generated(about: dict) -> bool:
    return about.get("generator") == GENERATOR


# ---------------------------------------------------------------------
# VALIDATION
# ---------------------------------------------------------------------

def compare_series(ref_t, ref_v, gen_t, gen_v, max_lag_minutes: float = MAX_LAG_MINUTES):
    """How a reference spreadsheet (ref) differs from a prediction (gen)
    over their overlap. The prediction is interpolated to the
    reference's times, never across a gap in it longer than two rows.

    Returns None if they share fewer than MIN_OVERLAP_ROWS rows,
    otherwise a dict with
      n, start, end               the overlap
      mean, rms, rms_demeaned,    ref - gen (m)
      max_abs, corr
      lag_min                     best L, in minutes, in ref(t) ~ gen(t - L):
                                  positive means the reference's times run
                                  late. The lag where a scaled and offset
                                  prediction fits the reference best, so a
                                  range difference does not move it. Found
                                  on a 1-minute scan, refined with a
                                  parabola through the three best points.
                                  None if the overlap is too short to scan
                                  +/-max_lag_minutes.
      rms_at_lag, ratio           at that lag: the mean-removed RMS, and
                                  the reference's range over the prediction's
                                  (regression slope)
      off_grid                    reference times not on a whole
                                  STEP_MINUTES mark
    """
    rx = _seconds(ref_t)
    rv = np.asarray(ref_v, dtype=float)
    order = np.argsort(_seconds(gen_t))
    gx = _seconds(gen_t)[order]
    gv = np.asarray(gen_v, dtype=float)[order]
    gap_s = 2 * STEP_MINUTES * 60

    g0 = _interp_no_gaps(rx, gx, gv, gap_s)
    ok = np.isfinite(g0) & np.isfinite(rv)
    if ok.sum() < MIN_OVERLAP_ROWS:
        return None
    d = rv[ok] - g0[ok]
    a, b = rv[ok], g0[ok]
    corr = (float(np.corrcoef(a, b)[0, 1]) if np.std(a) > 0 and np.std(b) > 0
            else float("nan"))
    ref_ok = [t for t, k in zip(ref_t, ok) if k]
    result = {
        "n": int(ok.sum()),
        "start": min(ref_ok),
        "end": max(ref_ok),
        "mean": float(np.mean(d)),
        "rms": float(np.sqrt(np.mean(d ** 2))),
        "rms_demeaned": float(np.std(d)),
        "max_abs": float(np.max(np.abs(d))),
        "corr": corr,
        "lag_min": None,
        "rms_at_lag": None,
        "ratio": None,
        "off_grid": int(sum(1 for t in ref_t
                            if t.second or t.microsecond or t.minute % STEP_MINUTES)),
    }

    # Lag scan, on the rows the prediction covers at every lag, so each
    # lag is scored on the same rows.
    max_lag_s = float(max_lag_minutes) * 60
    use = ok & (rx - max_lag_s >= gx[0]) & (rx + max_lag_s <= gx[-1])
    if use.sum() < MIN_OVERLAP_ROWS:
        return result
    lags = np.arange(-int(max_lag_minutes), int(max_lag_minutes) + 1, dtype=float)

    def misfit(lag_min):
        # what is left after the best linear fit, var(a) (1 - r^2): a
        # plain var(a - b) would trade the lag against the range, since
        # interpolating between rows slightly flattens the prediction
        g = _interp_no_gaps(rx[use] - lag_min * 60, gx, gv, gap_s)
        k = np.isfinite(g)
        a, b = rv[use][k], g[k]
        if k.sum() < MIN_OVERLAP_ROWS or np.std(a) == 0 or np.std(b) == 0:
            return np.inf
        return float(np.var(a) * (1 - np.corrcoef(a, b)[0, 1] ** 2))

    scores = np.array([misfit(L) for L in lags])
    i = int(np.argmin(scores))
    lag = float(lags[i])
    if 0 < i < len(lags) - 1 and np.isfinite(scores[i - 1:i + 2]).all():
        s0, s1, s2 = scores[i - 1], scores[i], scores[i + 1]
        curve = s0 - 2 * s1 + s2
        if curve > 0:
            lag += 0.5 * (s0 - s2) / curve
    g = _interp_no_gaps(rx[use] - lag * 60, gx, gv, gap_s)
    k = np.isfinite(g)
    a, b = rv[use][k], g[k]
    result["lag_min"] = lag
    result["rms_at_lag"] = float(np.std(a - b))
    result["ratio"] = (float(np.cov(a, b)[0, 1] / np.var(b, ddof=1))
                       if np.var(b) > 0 else float("nan"))
    return result


def is_clock_lag(lag_min: float) -> bool:
    """True for a lag that reads as a clock error: one 10-minute row,
    or 4-5 h of time zone (see VALIDATION)."""
    size = abs(lag_min)
    return size > LAG_SAME_MIN and (
        abs(size - STEP_MINUTES) <= LAG_SAME_MIN
        or any(abs(size - 60 * hours) <= TIME_ZONE_TOL_MIN for hours in TIME_ZONE_HOURS))


def read_lag(lag_min: float, same_model: bool = True) -> str:
    """What a best-fit lag most likely means (see VALIDATION). With
    same_model False (a NOAA prediction, or one of unknown source) the
    lag also holds a real timing difference, so it is not read as a
    clock error."""
    if not same_model:
        return ("not a clock error by itself: against another model or place (NOAA: the "
                "gauge's own delay) the lag is partly real; compare two spreadsheets' lags "
                "with each other")
    size = abs(lag_min)
    if size <= LAG_SAME_MIN:
        return "the two agree in time"
    if abs(size - STEP_MINUTES) <= LAG_SAME_MIN:
        return (f"one {STEP_MINUTES}-minute row -- the ambiguity merge_tide_files.py "
                f"describes")
    for hours in TIME_ZONE_HOURS:
        if abs(size - 60 * hours) <= TIME_ZONE_TOL_MIN:
            if lag_min < 0:
                return (f"a time-zone error: {hours} h is Eastern time "
                        f"({'EDT' if hours == 4 else 'EST'}) written as UTC")
            return (f"a time-zone error: the times are {hours} h late, e.g. converted "
                    f"to UTC twice")
    return ("a timing difference between the two predictions themselves (another model "
            "or place, e.g. a harbour), not a clock error")


def source_kind(about: dict) -> str:
    """"pytmd" or "noaa" for a prediction this script made, else ""."""
    return str(about.get("source_key") or "").split(":", 1)[0]


def format_report(name: str, column: str, label: str, res: dict | None,
                  kind: str = "pytmd") -> str:
    """One spreadsheet's comparison. `kind` is the prediction's source
    (source_kind): only against "pytmd" is the lag absolute, so only
    then does it say how many minutes to add to the spreadsheet."""
    same_model = kind == "pytmd"
    lines = [f"Spreadsheet   : {name}, column {column}",
             f"Prediction    : {label}"]
    if res is None:
        lines.append(f"Overlap       : fewer than {MIN_OVERLAP_ROWS} shared rows -- nothing "
                     f"to compare")
        return "\n".join(lines)
    lines += [
        f"Overlap       : {res['start']:%Y-%m-%d %H:%M} to {res['end']:%Y-%m-%d %H:%M} UTC "
        f"({res['n']} rows)",
        f"Mean diff     : {res['mean']:+.3f} m (spreadsheet minus prediction)",
        f"RMS diff      : {res['rms']:.3f} m; {res['rms_demeaned']:.3f} m with the mean "
        f"removed; largest {res['max_abs']:.3f} m",
        f"Correlation   : {res['corr']:.5f}",
    ]
    lag = res["lag_min"]
    if lag is None:
        lines.append("Best lag      : overlap too short to scan for a time lag")
    else:
        lag = 0.0 if abs(lag) < 0.05 else lag      # no "-0.0"
        shift = -STEP_MINUTES * int(round(lag / STEP_MINUTES))     # whole rows
        late = "late" if lag > 0 else "early"
        lines += [
            f"Best lag      : {lag:+.1f} min, RMS {res['rms_at_lag']:.3f} m there "
            f"(mean removed)",
        ]
        if same_model and is_clock_lag(lag):
            lines += [
                f"                the spreadsheet's times run {abs(lag):.1f} min {late}: add "
                f"{shift:+d} min to them to align",
                f"                (if this is the older of the two files: merge_tide_files.py "
                f"--shift-old-minutes {shift:+d})",
            ]
        elif not same_model:
            lines.append("                (includes the prediction's own timing, e.g. the "
                         "NOAA gauge's delay: not a shift to apply)")
        lines += [
            f"Height ratio  : {res['ratio']:.3f} (spreadsheet range / prediction range, "
            f"at the best lag)",
            f"Reading       : {read_lag(lag, same_model)}",
        ]
    lines.append(f"Off-grid rows : {res['off_grid']} (times not on a whole "
                 f"{STEP_MINUTES}-minute mark)")
    return "\n".join(lines)


def format_lag_differences(results) -> str:
    """The later spreadsheets' lags relative to the first one's, the
    first --reference taken as the older file. The prediction's own
    timing (a NOAA gauge's delay) cancels in the difference, so this
    holds whatever the source: it is the shift merge_tide_files.py
    --shift-old-minutes needs. `results` is [(name, compare_series
    result or None), ...]."""
    old_name, old_res = results[0]
    lines = [f"Between the spreadsheets ({old_name} taken as the older file):"]
    if old_res is None or old_res["lag_min"] is None:
        lines.append(f"  no lag for {old_name}, so nothing to compare")
        return "\n".join(lines)
    for name, res in results[1:]:
        if res is None or res["lag_min"] is None:
            lines.append(f"  {name}: no lag, so nothing to compare")
            continue
        shift = res["lag_min"] - old_res["lag_min"]
        rows = int(round(shift / STEP_MINUTES))
        lines.append(f"  {name} lag minus {old_name} lag: {shift:+.1f} min")
        if abs(shift - rows * STEP_MINUTES) > LAG_SAME_MIN:
            lines.append(f"  not a whole number of {STEP_MINUTES}-minute rows, so not a "
                         f"clock offset between the files: do not merge them with a shift "
                         f"(merge_tide_files.py would find no shared timestamps)")
            continue
        lines += [f"  merge_tide_files.py --old {old_name} --new {name} "
                  f"--shift-old-minutes {rows * STEP_MINUTES:+d}",
                  f"  Reading     : {old_name} against {name}: {read_lag(-shift)}"]
    return "\n".join(lines)


def _resolve(path_text: str) -> Path:
    """A path as given, else relative to the project (station_health
    resolves tide_model_file the same way)."""
    path = Path(path_text).expanduser()
    if not path.is_absolute() and not path.exists():
        path = PROJECT_DIR / path
    return path


def history_seed(cfg: dict, column: str):
    """(times, values, path, column) of the spreadsheet tide_model_file
    names, sorted by time, when the plot reads the height column
    `column` from it (tide_model_value_column); otherwise None, saying
    why."""
    path = _resolve(cfg["tide_model_file"])
    if cfg.get("tide_model_value_column") != column:
        print(f"  no rows taken from {path.name}: tide_model_value_column is "
              f"{cfg.get('tide_model_value_column')!r}, not {column!r}")
        return None
    try:
        times, values, _, _ = read_prediction(
            path, cfg.get("tide_model_time_column") or "time", column)
    except Exception as exc:
        print(f"  no rows taken from {path.name} ({exc})")
        return None
    if not times:
        return None
    order = sorted(range(len(times)), key=times.__getitem__)
    return [times[i] for i in order], values[order], path, column


def validate(args, cfg: dict, opts: dict) -> int:
    refs = args.reference or ([cfg["tide_model_file"]] if cfg.get("tide_model_file") else [])
    if not refs:
        print("No spreadsheet to compare with: give --reference, or set tide_model_file.")
        return 2
    ref_time_col = args.reference_time_col or cfg.get("tide_model_time_column") or "time"
    ref_value_col = args.reference_value_col or cfg.get("tide_model_value_column")
    loaded = []
    for r in refs:
        path = _resolve(r)
        times, values, _, column = read_prediction(path, ref_time_col, ref_value_col)
        if not times:
            print(f"{path.name}: no rows with a time and a height")
            return 1
        loaded.append((path, column, times, values))

    if args.generated:
        gen_path = _resolve(args.generated)
        gen_t, gen_v, about, _ = read_prediction(gen_path, opts["time_column"],
                                                 opts["value_column"])
        label = f"{gen_path.name}: {about.get('source', 'not made by ' + GENERATOR)}"
    else:
        pad = timedelta(minutes=args.max_lag_minutes + 60)
        gen_t = grid_times(min(min(t) for _, _, t, _ in loaded) - pad,
                           max(max(t) for _, _, t, _ in loaded) + pad)
        print(f"Predicting {len(gen_t)} times, {gen_t[0]:%Y-%m-%d} to {gen_t[-1]:%Y-%m-%d} "
              f"UTC ...")
        try:
            gen_v, about = predict(gen_t, opts)
        except Exception as exc:
            print(f"  Tide prediction failed: {exc}")
            return 1
        label = f"{about['source']}, {about['point']}"

    kind = source_kind(about)
    status = 0
    results = []
    for path, column, times, values in loaded:
        res = compare_series(times, values, gen_t, gen_v, args.max_lag_minutes)
        print()
        print("=" * 72)
        print(format_report(path.name, column, label, res, kind))
        if res is None:
            status = 1
        results.append((path.name, res))
    if len(results) > 1:
        print("=" * 72)
        print(format_lag_differences(results))
    print("=" * 72)
    return status


# ---------------------------------------------------------------------
# GENERATION
# ---------------------------------------------------------------------

def generate(args, cfg: dict, opts: dict) -> int:
    out = opts["output"]
    old = None
    if out.exists():
        try:
            old = read_prediction(out, opts["time_column"], None)
        except Exception as exc:
            old = None
            print(f"  could not read the existing {out.name} ({exc})")
        if (old is None or not is_generated(old[2])) and not args.force:
            print(f"  {out} exists and was not made by {GENERATOR} (no '{ABOUT_SHEET}' "
                  f"sheet), so it is left alone. Choose another --output, or --force.")
            return 2

    # "auto" never changes the height column the plot reads: falling
    # back to NOAA on a night pyTMD fails would rename it, and the plot
    # would lose its predicted tide altogether. Once tide_model_file is
    # this file, auto takes the source of tide_model_value_column's
    # column; otherwise it keeps the source of the file it already wrote.
    own_columns = {"pytmd": f"{opts['model']}_heightm",
                   "noaa": f"NOAA{opts['noaa_station']}_heightm"}
    configured = cfg.get("tide_model_file")
    plot_reads_out = bool(configured) and _resolve(configured).resolve() == out.resolve()
    if opts["source"] == "auto" and plot_reads_out and not opts["value_column"]:
        for kind, column in own_columns.items():
            if cfg.get("tide_model_value_column") == column:
                opts = dict(opts, source=kind)
                print(f"  auto: using {kind}, whose column {column} is "
                      f"tide_model_value_column")
                break
    if opts["source"] == "auto" and old is not None and is_generated(old[2]):
        kept = source_kind(old[2])
        if kept in ("pytmd", "noaa"):
            opts = dict(opts, source=kept)
            print(f"  auto: keeping {kept}, the source of the existing {out.name} "
                  f"(delete it, or give --source, to choose again)")

    if args.start:
        start = datetime.strptime(args.start, "%Y-%m-%d")
        end = start + timedelta(days=args.days)
    else:
        start, end = rolling_window(utcnow(), args.days_back, args.days)
    window_start = start

    # First run with history on: start from the spreadsheet the plot
    # reads now, so the months already published keep their tide (SAFETY).
    seed = None
    if opts["keep_history"] and not (old is not None and is_generated(old[2])) \
            and configured and not plot_reads_out:
        column = opts["value_column"] or own_columns.get(opts["source"], own_columns["pytmd"])
        seed = history_seed(cfg, column)
        if seed is not None and seed[0][-1] < start:
            start = seed[0][-1] + timedelta(minutes=STEP_MINUTES)
    times = grid_times(start, end)
    try:
        values, about = predict(times, opts)
    except Exception as exc:
        print(f"  Tide prediction failed: {exc}")
        print(f"  {out} is unchanged." if out.exists() else f"  {out} was not written.")
        return 1

    value_col = opts["value_column"] or about.pop("value_column")
    about.pop("value_column", None)
    if seed is not None and seed[3] != value_col:
        # auto fell back to NOAA: the spreadsheet's column is not this one
        print(f"  {seed[2].name}'s rows are not kept: its column {seed[3]} is not "
              f"{value_col}")
        keep = np.array([t >= window_start for t in times])
        times = [t for t, k in zip(times, keep) if k]
        values = values[keep]
        seed = None
    about.update({
        "time": "UTC. The times on the first sheet are naive; every one is UTC.",
        "interval": f"{STEP_MINUTES} minutes; each height is the instantaneous value at its time",
        "window": f"{times[0]:%Y-%m-%d %H:%M} to {times[-1]:%Y-%m-%d %H:%M} UTC",
        "generated_utc": f"{utcnow():%Y-%m-%d %H:%M:%S}",
    })

    if opts["keep_history"] and old is not None and is_generated(old[2]):
        old_t, old_v, old_about, old_col = old
        if old_about.get("source_key") != about["source_key"] or old_col != value_col:
            print(f"  previous {out.name} came from another source or column "
                  f"({old_about.get('source_key')}, {old_col}); its rows are not kept")
        else:
            earlier = [(t, v) for t, v in zip(old_t, old_v) if t < times[0]]
            if earlier:
                times = [t for t, _ in earlier] + times
                values = np.concatenate([[v for _, v in earlier], values])
                about["history_from"] = (old_about.get("history_from")
                                         or f"{earlier[0][0]:%Y-%m-%d %H:%M} UTC")
                if old_about.get("seeded_from"):
                    about["seeded_from"] = old_about["seeded_from"]
                print(f"  kept {len(earlier)} earlier row(s) from the previous file")
    elif seed is not None:
        seed_t, seed_v, seed_path, seed_col = seed
        earlier = [(t, v) for t, v in zip(seed_t, seed_v) if t < times[0]]
        if earlier:
            times = [t for t, _ in earlier] + times
            values = np.concatenate([[v for _, v in earlier], values])
            about["history_from"] = f"{earlier[0][0]:%Y-%m-%d %H:%M} UTC"
            about["seeded_from"] = (f"{seed_path.name}, column {seed_col}: its rows "
                                    f"{earlier[0][0]:%Y-%m-%d %H:%M} to "
                                    f"{earlier[-1][0]:%Y-%m-%d %H:%M} copied as they were; "
                                    f"this prediction from {times[len(earlier)]:%Y-%m-%d %H:%M}")
            print(f"  took {len(earlier)} earlier row(s) from {seed_path.name} "
                  f"(tide_model_file), so the months already published keep their tide")

    write_prediction(out, times, values, opts["time_column"], value_col, about)
    print(f"Wrote {out}")
    print(f"  {len(times)} rows, {times[0]:%Y-%m-%d %H:%M} to {times[-1]:%Y-%m-%d %H:%M} UTC, "
          f"every {STEP_MINUTES} minutes")
    print(f"  source: {about['source']} ({about['point']})")
    print(f"  columns: {opts['time_column']}, {value_col}")
    if not plot_reads_out:
        print(f"  The 7-day plot reads tide_model_file ({configured or 'unset'}); it uses this "
              f"file only once tide_model_file is \"{out.resolve()}\" and "
              f"tide_model_value_column is \"{value_col}\".")
    elif cfg.get("tide_model_value_column") != value_col:
        print(f"  WARNING: tide_model_file is this file, but tide_model_value_column is "
              f"{cfg.get('tide_model_value_column')!r}, not {value_col!r}: the 7-day plot "
              f"finds no predicted tide in it.")
    return 0


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--source", choices=("auto", "pytmd", "noaa"), default=None,
                   help="default: tide_prediction.source in station.json, else auto")
    p.add_argument("--output", default=None,
                   help="default: tide_prediction.file in station.json, else "
                        "tides/tide_prediction.xlsx in the project; a relative path is "
                        "taken from the project directory, not the current one")
    p.add_argument("--days", type=int, default=DAYS,
                   help=f"length of the prediction (default {DAYS})")
    p.add_argument("--days-back", type=int, default=DAYS_BACK,
                   help=f"days before today (UTC) it starts (default {DAYS_BACK})")
    p.add_argument("--start", default=None,
                   help="YYYY-MM-DD (UTC) to start on instead of the rolling window")
    p.add_argument("--lat", type=float, default=None,
                   help="default: tide_prediction.point, else station.json latitude")
    p.add_argument("--lon", type=float, default=None,
                   help="default: tide_prediction.point, else station.json longitude")
    p.add_argument("--pytmd-directory", default=None,
                   help="directory holding EOT20/ocean_tides/ (default: pyTMD's own)")
    p.add_argument("--model", default=None, help=f"pyTMD model (default {DEFAULT_MODEL})")
    p.add_argument("--extrapolate-km", type=float, default=None,
                   help=f"fill a land cell from ocean within this distance "
                        f"(default {EXTRAPOLATE_KM:g})")
    p.add_argument("--noaa-station", default=None,
                   help=f"CO-OPS harmonic station (default {DEFAULT_NOAA_STATION}, Chatham)")
    p.add_argument("--noaa-time-offset-minutes", type=float, default=None,
                   help="the site's tide comes this much later than the station's (default 0)")
    p.add_argument("--noaa-height-ratio", type=float, default=None,
                   help="the site's tide is this many times the station's (default 1)")
    p.add_argument("--time-col", default=None,
                   help="time column name (default: tide_model_time_column, else 'time')")
    p.add_argument("--value-col", default=None,
                   help="height column name (default EOT20_heightm, or NOAA<station>_heightm)")
    p.add_argument("--keep-history", action="store_true", default=None,
                   help="keep the earlier rows of this script's previous file (default: "
                        "tide_prediction.keep_history, else on)")
    p.add_argument("--no-keep-history", dest="keep_history", action="store_false",
                   help="keep only the new window")
    p.add_argument("--force", action="store_true",
                   help="overwrite --output even if this script did not make it")
    v = p.add_argument_group("validation")
    v.add_argument("--validate", action="store_true",
                   help="compare with the existing spreadsheet(s) instead of writing a file")
    v.add_argument("--reference", nargs="+", default=None,
                   help="spreadsheet(s) to compare with (default: tide_model_file)")
    v.add_argument("--reference-time-col", default=None)
    v.add_argument("--reference-value-col", default=None,
                   help="default: tide_model_value_column, else the first height column")
    v.add_argument("--generated", default=None,
                   help="compare this already-made prediction instead of making one")
    v.add_argument("--max-lag-minutes", type=float, default=MAX_LAG_MINUTES)
    args = p.parse_args(argv)

    cfg = load_station()
    opts = resolve_options(args, cfg)
    if args.validate:
        return validate(args, cfg, opts)
    return generate(args, cfg, opts)


if __name__ == "__main__":
    raise SystemExit(main())
